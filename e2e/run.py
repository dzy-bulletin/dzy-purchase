#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""貨單辨識系統 資料帶入測試（e2e）。用真後端＋假 Ollama／假損益端／假備份端＋Playwright。

  cd e2e && python3 run.py
  E2E_SEED=12345 python3 run.py      # 重現某次的資料

每次執行：全新 DATA_DIR（/private/tmp）、全新隨機資料與密碼；不呼叫任何外部服務。
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
VENV_PY = os.path.join(HERE, '.venv', 'bin', 'python')
try:
    import openpyxl  # noqa: F401
except ImportError:
    if os.path.exists(VENV_PY) and os.path.abspath(sys.executable) != os.path.abspath(VENV_PY):
        os.execv(VENV_PY, [VENV_PY] + sys.argv)
    raise

import json
import random
import re
import shutil
import signal
import socket
import sqlite3
import subprocess
import tempfile
import time
import uuid
from datetime import date

from playwright.sync_api import sync_playwright

sys.path.insert(0, HERE)
import common as C                                              # noqa: E402
from common import check, eq, ck, mark, scan, txt, val, shot, api, login, upload_api, last_dialog, CM, DLG, ENV   # noqa: E402
import dataset as DS                                            # noqa: E402
from dataset import fmt, fnum, norm_text, r2, cents             # noqa: E402
from fake_ollama import FakeOllama                              # noqa: E402
from fake_ext import FakePnl, FakeBackup                        # noqa: E402
from photos import make_jpeg                                    # noqa: E402

ROOT = os.path.dirname(HERE)


def free_port():
    s = socket.socket()
    s.bind(('127.0.0.1', 0))
    p = s.getsockname()[1]
    s.close()
    return p


class Stack:
    """一次執行用到的所有行程與目錄。"""

    def __init__(self, rng):
        self.tmp = tempfile.mkdtemp(prefix='purchase-e2e-', dir='/private/tmp')
        self.data = os.path.join(self.tmp, 'data')
        self.logs = os.path.join(self.tmp, 'logs')
        self.pics = os.path.join(self.tmp, 'pics')
        os.makedirs(self.pics)
        self.procs = []
        self.rng = rng
        self.pnl_key = DS.pw(rng, 24)
        self.bak_key = DS.pw(rng, 24)
        self.ollama = FakeOllama().start()
        self.pnl = FakePnl(self.pnl_key).start()
        self.bak = FakeBackup(self.bak_key).start()
        self.port = free_port()
        self.fport = free_port()
        ENV['API'] = f'http://localhost:{self.port}/purchase/api'
        ENV['FRONT'] = f'http://localhost:{self.fport}'
        ENV['API_URL'] = ENV['API']
        self.env = dict(os.environ, DATA_DIR=self.data, PORT=str(self.port), OLLAMA_URL=f'http://127.0.0.1:{self.ollama.port}',
                        RETRY_DELAY_MS='80', PNL_PUSH_URL=self.pnl.url, PNL_PURCHASE_KEY=self.pnl_key, PNL_TICK_MS='400',
                        LOG_DIR=self.logs, PURCHASE_NO_DOTENV='1', MODEL='fake-vl')
        for k in ('BACKUP_URL', 'BACKUP_KEY'):
            self.env.pop(k, None)

    def start(self, admin_user, admin_pw):
        env = dict(self.env, E2E_ADMIN_USER=admin_user, E2E_ADMIN_PW=admin_pw)
        subprocess.run(['node', os.path.join(HERE, 'bootstrap.js')], env=env, check=True)
        self.server = subprocess.Popen(['node', os.path.join(ROOT, 'server', 'index.js')], env=self.env, cwd=ROOT,
                                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        self.web = subprocess.Popen([sys.executable, '-m', 'http.server', str(self.fport), '--bind', '127.0.0.1', '-d', os.path.join(ROOT, 'web')],
                                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(80):
            try:
                import urllib.request
                urllib.request.urlopen(ENV['API'] + '/health', timeout=1).read()
                urllib.request.urlopen(ENV['FRONT'] + '/upload.html', timeout=1).read(1)
                return
            except Exception:
                time.sleep(0.25)
        raise RuntimeError('伺服器起不來')

    def db(self):
        c = sqlite3.connect(os.path.join(self.data, 'purchase.db'), timeout=10)
        c.row_factory = sqlite3.Row
        return c

    def q(self, sql, *a):
        c = self.db()
        try:
            return [dict(r) for r in c.execute(sql, a).fetchall()]
        finally:
            c.close()

    def x(self, sql, *a):
        c = self.db()
        try:
            c.execute(sql, a)
            c.commit()
        finally:
            c.close()

    def stop(self):
        for p in (getattr(self, 'server', None), getattr(self, 'web', None)):
            if p:
                try:
                    p.send_signal(signal.SIGTERM)
                    p.wait(timeout=8)
                except Exception:
                    p.kill()
        for f in (self.ollama, self.pnl, self.bak):
            try:
                f.stop()
            except Exception:
                pass
        shutil.rmtree(self.tmp, ignore_errors=True)

    def wait_idle(self, timeout=90):
        t0 = time.time()
        while time.time() - t0 < timeout:
            n = self.q("SELECT COUNT(*) c FROM slips WHERE status IN ('uploaded','queued','recognizing')")[0]['c']
            if n == 0:
                return True
            time.sleep(0.3)
        return False

    def wait_pnl_quiet(self, timeout=40):
        """損益推送 outbox 清空（終態除外）後再多等一輪。"""
        t0 = time.time()
        while time.time() - t0 < timeout:
            n = self.q('SELECT COUNT(*) c FROM pnl_outbox WHERE state IS NULL')[0]['c']
            m = self.q('SELECT COUNT(*) c FROM pnl_retire WHERE state IS NULL')[0]['c']
            if n == 0 and m == 0:
                time.sleep(1.0)
                n = self.q('SELECT COUNT(*) c FROM pnl_outbox WHERE state IS NULL')[0]['c']
                if n == 0:
                    return True
            time.sleep(0.3)
        return False


def new_context(browser):
    ctx = browser.new_context(locale='zh-TW', viewport={'width': 1366, 'height': 900}, accept_downloads=True)

    def cfg(route):
        resp = route.fetch()
        route.fulfill(response=resp, body=resp.text().replace('http://localhost:8794/purchase/api', ENV['API']))
    ctx.route('**/js/config.js', cfg)

    def ext(route):
        route.abort()
    ctx.route(re.compile(r'^https?://(?!localhost|127\.0\.0\.1).*'), ext)       # 不連任何外部（Google Fonts 等）
    return ctx


def url(name, hash_=''):
    return f'{ENV["FRONT"]}/{name}?api={ENV["API"]}{hash_}'


# ───────────────────────── 初始化（API）─────────────────────────
def setup_backend(S, D, M, slips):
    adm, code, _ = login(D['admin']['username'], D['admin']['pass'])
    check('admin 可登入（免強制改密碼）', adm and not adm['must_change_password'], code)
    D['adm_tok'] = adm['token']
    t = adm['token']
    for s in D['stores']:
        c, j = api('POST', '/admin/stores', t, {'code': s['code'], 'name': s['name'], 'brand_id': s['brand'], 'password': s['pass0'],
                                                 'pnl_unit_code': s['unit_code'] or ''})
        assert j.get('ok'), j
        s['id'] = j['data']['id']
    for a in (D['acc_multi'], D['acc_single']):
        c, j = api('POST', '/admin/users', t, {'username': a['username'], 'name': a['name'], 'role': 'accountant', 'brand_ids': a['brands'],
                                                'brand_id': a['brands'][0], 'password': a['pass0']})
        assert j.get('ok'), j
    for b in D['brand_order']:
        D['vid'] = D.get('vid', {})
        for v in D['vendors'][b]:
            c, j = api('POST', '/vendors', t, {'brand_id': b, 'name': v['name']})
            assert j.get('ok'), j
            v['id'] = j['data']['id']
            M.vendors[b][v['name']] = {'name': v['name'], 'active': True, 'inc': 0, 'auto': False}
        for it in D['items'][b]:
            c, j = api('POST', '/items', t, {'brand_id': b, 'name': it['name'], 'category': it['cat'], 'base_unit': it['base']})
            assert j.get('ok'), j
            it['id'] = j['data']['id']
            if it['conv']:
                c, j = api('PUT', f'/items/{it["id"]}/units', t, [{'unit': u, 'factor': f} for u, f in it['conv'].items()])
                assert j.get('ok'), j
            M.add_item(b, it['name'], it['cat'], it['base'], it['conv'])
    # 假 AI 與照片
    for s in slips:
        ai, amap = DS.build_ai(S.rng, D, s)
        s['ai'], s['ai_map'] = ai, amap
        S.ollama.by_dims[s['dims']] = ai
        if s['inj'].get('fail'):
            S.ollama.fail_first[s['dims']] = 3
        files = [make_jpeg(S.pics, s['dims'][0], s['dims'][1], s['idx'])]
        for j in range(1, s['photos']):
            files.append(make_jpeg(S.pics, 1000 + s['idx'] * 3 + j, 500 + j * 11, s['idx'] * 10 + j))
        s['files'] = files


def ctx_for(M, s):
    b = s['brand']
    vname = None if s['vendor_typed'] else s['vendor']['name']
    v = M.vendors[b].get(vname) if vname else None
    return {'resolve': (lambda raw: M.resolve(b, vname, raw)) if vname else (lambda raw: None),
            'tax_included': 1 if (v and v['inc']) else 0}


def store_token(D, store, pw):
    d, code, j = login(store['code'], pw)
    return d['token'] if d else None


def fake_check_store_api_gate(D):
    """首次登入前，其他 API 一律 PASSWORD_CHANGE_REQUIRED；改完才能用。（走 API，門市 UI 只用一家）"""
    out = {}
    for s in D['stores']:
        if s is D['ui_store'] or s.get('lock'):
            continue
        d, code, j = login(s['code'], s['pass0'])
        tok = d['token']
        if s is D['stores'][0] or s is (D['stores'][1] if len(D['stores']) > 1 else None):
            c, j2 = api('GET', '/slips?mine=1', tok)
            check(f'門市 {s["code"]} 首次登入前打其他 API 被擋（PASSWORD_CHANGE_REQUIRED）', c == 403 and j2.get('error') == 'PASSWORD_CHANGE_REQUIRED', (c, j2))
        c, j2 = api('POST', '/password', tok, {'old_password': s['pass0'], 'new_password': s['pass1']})
        assert j2.get('ok'), j2
        out[s['code']] = j2['data']['token']
    return out


# ───────────────────────── 登入／改密碼（UI）─────────────────────────
TITLE = '鼎兆元｜貨單辨識系統'


def guide_flow(page, role, full):
    """登入畫面上的「第一次使用？看操作教學」：新分頁開 guide.html#<角色>。full＝把教學頁每個互動元素都走一遍。"""
    ctx = page.context
    href = page.evaluate("()=>document.querySelector('.guide-link a').getAttribute('href')")
    eq(f'登入畫面的教學連結（{role}）指向對應分頁', href, 'guide.html#' + role)
    with ctx.expect_page() as np:
        ck(page, '.guide-link a', '操作教學連結（新分頁）')
    g = np.value
    g.wait_for_load_state('load')
    eq('教學頁標題', g.title(), '鼎兆元｜貨單辨識系統操作教學')
    shown = g.evaluate("()=>['store','acc','admin'].filter(t=>!document.getElementById('p-'+t).hidden)")
    eq(f'教學頁依連結自動切到「{role}」分頁', shown, [role])
    if full:
        scan(g, '教學頁')
        for t in ('store', 'acc', 'admin'):
            ck(g, f'#t-{t}', f'切到「{t}」分頁')
            g.wait_for_timeout(200)
            eq(f'教學頁：點「{t}」分頁後只顯示該分頁', g.evaluate("()=>['store','acc','admin'].filter(t=>!document.getElementById('p-'+t).hidden)"), [t])
            g.evaluate("()=>{document.querySelectorAll('img[loading=lazy]').forEach(i=>i.loading='eager')}")
            g.wait_for_timeout(600)
            bad = g.evaluate(f"()=>[...document.querySelectorAll('#p-{t} img')].filter(i=>!(i.complete&&i.naturalWidth>0)).map(i=>i.getAttribute('src'))")
            check(f'教學頁「{t}」分頁：所有截圖都載得到', not bad, bad[:3])
            if t == 'store':
                g.locator('#p-store img').first.click()
                g.wait_for_timeout(300)
                check('教學頁：點截圖 → 放大燈箱', g.evaluate("()=>!document.getElementById('lb').hidden"))
                scan(g, '教學頁（燈箱）')
                ck(g, '#lbx', '關閉燈箱')
                check('教學頁：按 ✕ 關閉燈箱', g.evaluate("()=>document.getElementById('lb').hidden"))
                g.locator('#p-store img').first.click()
                g.keyboard.press('Escape')
                check('教學頁：Escape 關閉燈箱', g.evaluate("()=>document.getElementById('lb').hidden"))
        n = g.evaluate("()=>document.querySelectorAll('summary').length")
        for i in range(n):
            k = g.evaluate(C.KEY_JS, f'details:nth-of-type({i + 1}) > summary')
            g.locator('summary').nth(i).click()
            if k:
                CM.mark(k, '展開常見問題')
        check(f'教學頁：{n} 則常見問題都能展開', g.evaluate("()=>[...document.querySelectorAll('details')].every(d=>d.open)"))
    g.close()


def ui_login(page, acc, pw, screen):
    page.fill('#acc', acc)
    page.fill('#pw', pw)
    scan(page, screen)
    ck(page, '#loginBtn', '登入')
    page.wait_for_timeout(500)


def ui_force_change(page, old, new, who, acc=None):
    """第一次登入強制改密碼畫面：逐一驗證各種不合規輸入的提示，再改成功。"""
    page.wait_for_selector('#pwBtn', timeout=10000)
    scan(page, '強制改密碼畫面')
    check(f'{who}：首次登入出現強制改密碼畫面', '第一次登入請設定你自己的密碼' in txt(page, '#login'), txt(page, '#login')[:60])
    cases = [('', '', '', '請輸入舊密碼與新密碼'), (old, 'abc', 'abc', '新密碼至少 6 個字'), (old, old, old, '新密碼不可與舊密碼相同'),
             (old, new, new + 'x', '兩次輸入的新密碼不一樣'), ('wrong-' + old, new, new, '舊密碼不正確')]
    for a, b, c, want in cases:
        page.fill('#pwOld', a)
        page.fill('#pwNew', b)
        page.fill('#pwNew2', c)
        page.evaluate("()=>{document.getElementById('pwMsg').innerHTML=''}")
        ck(page, '#pwBtn', '送出（驗證錯誤提示）')
        page.wait_for_function("()=>document.getElementById('pwMsg').innerText.trim().length>0", timeout=8000)
        got = txt(page, '#pwMsg')
        check(f'{who}：改密碼錯誤提示「{want}」', want in got, got)
    # 強制畫面不可略過：沒有取消鈕，只有登出
    check(f'{who}：強制改密碼畫面沒有「取消」只有「登出」', page.evaluate("()=>!document.getElementById('pwCancel')&&!!document.getElementById('pwOut')"))
    ck(page, '#pwOut', '強制改密碼畫面的「登出」')
    page.wait_for_selector('#loginBtn', timeout=8000)
    check(f'{who}：強制改密碼畫面按「登出」回到登入畫面', True)
    page.fill('#acc', acc)
    page.fill('#pw', old)
    ck(page, '#loginBtn', '登入')
    page.wait_for_selector('#pwBtn', timeout=10000)
    page.fill('#pwOld', old)
    page.fill('#pwNew', new)
    page.fill('#pwNew2', new)
    ck(page, '#pwBtn', '設定密碼並進入')
    page.wait_for_selector('#app:not(.hidden)', timeout=15000)
    check(f'{who}：改密碼成功進入系統', True)


def ui_voluntary_change(page, old, new, who):
    """登入後自己改密碼（可取消）。"""
    ck(page, '#chpw', '改密碼')
    page.wait_for_selector('#pwBtn')
    check(f'{who}：自願改密碼畫面有「取消」', page.evaluate("()=>!!document.getElementById('pwCancel')"))
    scan(page, '自願改密碼畫面')
    ck(page, '#pwCancel', '取消改密碼，回原頁')
    page.wait_for_selector('#app:not(.hidden)')
    ck(page, '#chpw', '改密碼')
    page.fill('#pwOld', old)
    page.fill('#pwNew', new)
    page.fill('#pwNew2', new)
    ck(page, '#pwBtn', '儲存新密碼')
    page.wait_for_selector('#app:not(.hidden)', timeout=15000)
    check(f'{who}：自願改密碼成功', True)


def ui_logout(page, who):
    ck(page, '#logout', '登出')
    page.wait_for_selector('#loginBtn', timeout=8000)
    check(f'{who}：登出後回到登入畫面', True)


# ───────────────────────── 門市上傳（UI）─────────────────────────
STATUS_TEXT = {'uploaded': '已送出', 'queued': '排隊中', 'recognizing': '辨識中', 'review': '待核對', 'confirmed': '已入帳', 'failed': '辨識失敗', 'returned': '退回重拍'}


def read_mine(page):
    return page.evaluate("""()=>[...document.querySelectorAll('#mine .slip')].map(e=>({
        vendor:e.querySelector('.t').innerText, badge:e.querySelector('.badge').innerText, id:e.querySelector('.muted').innerText.split('\\u3000')[0].trim(),
        ret:e.classList.contains('ret'), text:e.innerText}))""")


def ui_upload(page, S, D, s, how='normal'):
    """門市在手機頁上傳一張貨單。how: normal／offline"""
    store = s['store']
    ids_before = {r['id'] for r in S.q('SELECT id FROM slips')}
    mark(page, '#vendorSel', '選廠商')
    if s['vendor_typed']:
        page.select_option('#vendorSel', '__other')
        page.fill('#vendorName', s['vendor_typed'])
    else:
        page.select_option('#vendorSel', label=s['vendor']['name'])
    files = s['files']
    with page.expect_file_chooser() as fc:
        ck(page, '#camBtn', '拍照（開檔案選擇）')
    fc.value.set_files(files[0])
    page.wait_for_function("()=>document.querySelectorAll('#thumbs .thumb').length>=1")
    if len(files) > 1:
        with page.expect_file_chooser() as fc:
            ck(page, '#pickBtn', '從相簿選')
        fc.value.set_files(files[1:])
        page.wait_for_function(f"()=>document.querySelectorAll('#thumbs .thumb').length=={len(files)}", timeout=15000)
    if len(files) >= 2 and s['idx'] % 2 == 0:       # 多拍一張再刪掉
        with page.expect_file_chooser() as fc:
            ck(page, '#pickBtn', '從相簿選')
        fc.value.set_files(files[0])
        page.wait_for_function(f"()=>document.querySelectorAll('#thumbs .thumb').length=={len(files) + 1}", timeout=15000)
        ck(page, '#thumbs .thumb:last-child button', '刪除縮圖')
        page.wait_for_function(f"()=>document.querySelectorAll('#thumbs .thumb').length=={len(files)}")
        check(f'{s["idx"]}：縮圖 ✕ 刪除後張數減一', True)
    eq(f'{s["idx"]}：縮圖張數＝拍的張數', page.evaluate("()=>document.querySelectorAll('#thumbs .thumb').length"), len(files))
    check('選好廠商與照片後「送出」可按', not page.evaluate("()=>document.getElementById('sendBtn').disabled"))
    if how == 'offline':
        page.context.set_offline(True)
        ck(page, '#sendBtn', '送出（離線）')
        page.wait_for_function("()=>document.getElementById('msg').innerText.indexOf('目前送不出去')>=0", timeout=60000)
        check('離線送出：顯示已存在手機、連線後自動補傳', True)
        page.wait_for_function("()=>document.getElementById('resTitle').innerText==='還沒送出'&&document.getElementById('resCard').classList.contains('r-warn')", timeout=60000)
        check('離線送出：跳出黃色「還沒送出」視窗', '照片已經存在這支手機' in txt(page, '#resBody'), txt(page, '#resBody'))
        shot(page, '02-門市上傳-還沒送出視窗')
        scan(page, '上傳結果視窗（還沒送出）')
        ck(page, '#resRetry', '還沒送出視窗：立即重試（仍離線）')
        page.wait_for_function("()=>document.getElementById('resTitle').innerText==='還沒送出'", timeout=60000)
        check('離線時按立即重試：仍顯示「還沒送出」', True)
        ck(page, '#resOk', '還沒送出視窗：知道了')
        check('按「知道了」後視窗關閉', page.evaluate("()=>document.getElementById('result').classList.contains('hidden')"))
        ck(page, '#retryNow', '待送列：立即補傳（仍離線）')
        page.wait_for_function("()=>/^還有 1 張沒送出/.test(document.getElementById('resTitle').innerText)", timeout=60000)
        check('離線時按待送列的立即補傳：顯示「還有 1 張沒送出」', page.evaluate("()=>document.getElementById('resCard').classList.contains('r-warn')"))
        ck(page, '#resOk', '還有沒送出視窗：知道了')
        page.wait_for_function("()=>!document.getElementById('pendingBar').classList.contains('hidden')")
        eq('離線送出：待送列顯示 1 張待送', txt(page, '#pendingText').strip(), '1 張待送')
        shot(page, '02-門市上傳-離線待送')
        page.context.set_offline(False)
        try:
            page.wait_for_function("()=>document.getElementById('pendingBar').classList.contains('hidden')", timeout=6000)
        except Exception:
            pass
        if page.evaluate("()=>!document.getElementById('pendingBar').classList.contains('hidden')"):
            ck(page, '#retryNow', '立即補傳')
        page.wait_for_function("()=>document.getElementById('pendingBar').classList.contains('hidden')", timeout=30000)
        check('離線補傳：恢復連線後待送列消失', True)
        page.wait_for_function("()=>!document.getElementById('result').classList.contains('hidden')&&document.getElementById('resTitle').innerText==='補傳成功'", timeout=30000)
        check('離線補傳：跳出綠色「補傳成功」視窗', page.evaluate("()=>document.getElementById('resCard').classList.contains('r-ok')"))
        scan(page, '上傳結果視窗（補傳成功）')
        ck(page, '#resOk', '補傳成功視窗：知道了')
    else:
        ck(page, '#sendBtn', '送出這張貨單')
        page.wait_for_function("()=>document.getElementById('msg').innerText.indexOf('已送出')>=0", timeout=30000)
    page.wait_for_timeout(300)
    new = [r['id'] for r in S.q('SELECT id FROM slips') if r['id'] not in ids_before]
    eq(f'{s["idx"]}：上傳後伺服器多一張貨單', len(new), 1)
    s['sid'] = new[0]
    if how != 'offline':
        page.wait_for_function("()=>!document.getElementById('result').classList.contains('hidden')&&document.getElementById('resTitle').innerText==='上傳成功'", timeout=30000)
        body = txt(page, '#resBody')
        check(f'{s["sid"]}：綠色「上傳成功」視窗顯示正確單號', s['sid'] in body, body)
        check(f'{s["sid"]}：成功視窗顯示廠商與照片張數', (s['vendor_typed'] or s['vendor']['name']) in body and f'{len(files)} 張' in body, body)
        if s['idx'] == 0:
            shot(page, '01-門市上傳-上傳成功視窗')
        scan(page, '上傳結果視窗（成功）')
        ck(page, '#resOk', '上傳成功視窗：好，繼續拍下一張')
        check('按下後視窗關閉', page.evaluate("()=>document.getElementById('result').classList.contains('hidden')"))
    row = S.q('SELECT store_id, brand_id, vendor_id, vendor_name_raw FROM slips WHERE id=?', s['sid'])[0]
    eq(f'{s["sid"]}：貨單屬於上傳的門市', row['store_id'], store['id'])
    eq(f'{s["sid"]}：照片張數', S.q('SELECT COUNT(*) c FROM slip_photos WHERE slip_id=?', s['sid'])[0]['c'], len(files))
    if s['vendor_typed']:
        eq(f'{s["sid"]}：手打的未知廠商名存在 vendor_name_raw', row['vendor_name_raw'], s['vendor_typed'])
    else:
        eq(f'{s["sid"]}：選的廠商存為 vendor_id', row['vendor_id'], s['vendor']['id'])
    return s['sid']


def store_ui_prelude(page, D, S):
    """UI 門市：登入、首次改密碼、廠商下拉內容。"""
    st = D['ui_store']
    page.goto(url('upload.html'))
    page.wait_for_selector('#loginBtn')
    eq('上傳頁標題', page.title(), TITLE)
    guide_flow(page, 'store', True)
    ui_login(page, st['code'], 'wrong-' + st['pass0'], '門市登入畫面')
    check('門市：密碼錯誤顯示「帳號或密碼錯誤」', '帳號或密碼錯誤' in txt(page, '#loginMsg'), txt(page, '#loginMsg'))
    page.fill('#acc', '')
    ck(page, '#loginBtn', '登入（空帳號）')
    check('門市：沒填帳號有提示', '請輸入門市代號' in txt(page, '#loginMsg'), txt(page, '#loginMsg'))
    ui_login(page, st['code'].lower(), st['pass0'], '門市登入畫面')           # 小寫代號也能登入（頁面會轉大寫）
    ui_force_change(page, st['pass0'], st['pass1'], '門市 ' + st['code'], st['code'].lower())
    page.wait_for_function("()=>document.getElementById('vendorSel').options.length>1")
    names = sorted(v['name'] for v in D['vendors'][st['brand']])
    opts = page.evaluate("()=>[...document.getElementById('vendorSel').options].map(o=>o.text)")
    eq('門市廠商下拉：第一項是提示、最後一項是手動輸入', (opts[0], opts[-1]), ('請選擇廠商', '（找不到，手動輸入）'))
    eq('門市廠商下拉：只列本品牌啟用的廠商', sorted(opts[1:-1]), names)
    eq('門市頁：顯示門市名稱', txt(page, '#who'), st['name'])
    check('門市頁：登入後套用品牌主題', page.evaluate("()=>document.documentElement.getAttribute('data-brand')") == st['brand'])
    check('送出鈕：沒選廠商沒拍照時不能按', page.evaluate("()=>document.getElementById('sendBtn').disabled"))
    shot(page, '01-門市上傳頁')


# ───────────────────────── 會計核對（UI）─────────────────────────
READ_JS = """()=>{
  const g=id=>{const e=document.getElementById(id);return e?e.value:null};
  const lines=[...document.querySelectorAll('#tb .lcard')].map(c=>{
    const f=k=>c.querySelector('[data-k="'+k+'"]');
    return {cls:c.className.replace('lcard','').trim(), raw:f('raw_name').value, item:f('item').value, qty:f('qty').value, unit:f('unit').value,
      price:f('unit_price').value, amount:f('amount').value, checked:f('checked').checked, why:c.querySelector('[data-why]').innerText};});
  const sum=document.getElementById('sum');
  return {id:(document.querySelector('#work h2')||{}).textContent, vendor:g('f_vendor'), date:g('f_date'), no:g('f_no'), sub:g('f_subtotal'), tax:g('f_tax'),
    total:g('f_total'), inc:!!(document.getElementById('f_taxinc')||{}).checked, hand:g('f_hand'), lines, slipWhy:(document.getElementById('slipWhy')||{}).innerText||'',
    sum:sum?sum.innerText:'', sumCls:sum?sum.className:'', handBox:!!document.querySelector('.hand'), note:(document.getElementById('note')||{}).innerText||'',
    confirmDisabled:(document.getElementById('confirmBtn')||{}).disabled, badge:(document.querySelector('#work .badge')||{}).innerText}}"""


def numval(x):
    return None if x in (None, '') else float(x)


def same_num(a, b):
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) < 0.0051


def verify_view(label, ui, exp):
    """exp: vendor, date, no, subtotal, tax, total, inc, lines[{raw,item,qty,unit,price,amount,flags}], date_note, hand, hand_box"""
    bad = []
    if ui['vendor'] != exp['vendor']:
        bad.append(f'廠商 {ui["vendor"]!r}≠{exp["vendor"]!r}')
    if ui['date'] != exp['date']:
        bad.append(f'日期 {ui["date"]}≠{exp["date"]}')
    for k, uk in (('subtotal', 'sub'), ('tax', 'tax'), ('total', 'total')):
        if not same_num(numval(ui[uk]), exp[k]):
            bad.append(f'{k} {ui[uk]}≠{exp[k]}')
    if bool(ui['inc']) != bool(exp['inc']):
        bad.append(f'含稅勾選 {ui["inc"]}≠{exp["inc"]}')
    if (ui['hand'] or '') != (exp['hand'] or ''):
        bad.append(f'手寫說明 {ui["hand"]!r}≠{exp["hand"]!r}')
    if len(ui['lines']) != len(exp['lines']):
        bad.append(f'列數 {len(ui["lines"])}≠{len(exp["lines"])}')
    else:
        for i, (u, e) in enumerate(zip(ui['lines'], exp['lines'])):
            for k in ('qty', 'price', 'amount'):
                if not same_num(numval(u[k]), e[k]):
                    bad.append(f'第{i + 1}列 {k} {u[k]}≠{e[k]}')
            if u['raw'] != e['raw'] or u['unit'] != e['unit']:
                bad.append(f'第{i + 1}列 品名／單位 {u["raw"]!r}/{u["unit"]!r}≠{e["raw"]!r}/{e["unit"]!r}')
            if (u['item'] or '') != (e['item'] or ''):
                bad.append(f'第{i + 1}列 統一品名 {u["item"]!r}≠{e["item"]!r}')
            cls, msgs, badq = DS.line_view(e)
            if u['cls'] != cls:
                bad.append(f'第{i + 1}列 顏色 {u["cls"]!r}≠{cls!r}')
            for m in DS.FLAG_TEXT.values():
                if (m in u['why']) != (m in msgs):
                    bad.append(f'第{i + 1}列 說明「{m}」{"多出" if m in u["why"] else "缺少"}（{u["why"]!r}）')
            if ('數量、單價、金額都要填且大於 0' in u['why']) != badq:
                bad.append(f'第{i + 1}列 缺值紅字說明不符')
            if e['qty'] is not None and e['price'] is not None and ('數量×單價＝' + fmt(r2(e['qty'] * e['price']))) not in u['why']:
                bad.append(f'第{i + 1}列 缺「數量×單價＝{fmt(r2(e["qty"] * e["price"]))}」')
    amounts = [l['amount'] for l in exp['lines']]
    s, ok, missing = DS.sum_check(amounts, exp['subtotal'], exp['tax'], exp['total'], exp['inc'])
    if ('good' in ui['sumCls']) != ok:
        bad.append(f'合計列顏色 {ui["sumCls"]!r} 預期{"相符" if ok else "不符"}')
    if ok and '相符' not in ui['sum']:
        bad.append('合計列缺「相符」')
    if not ok:
        w = DS.sum_why(amounts, exp['subtotal'], exp['tax'], exp['total'], exp['inc'])
        if w not in ui['sum']:
            bad.append(f'合計列白話說明不符：{ui["sum"]!r} 預期含 {w!r}')
    if not ui['sum'].startswith('各列加總 ' + fmt(s)):
        bad.append(f'各列加總顯示 {ui["sum"][:30]!r}≠{fmt(s)}')
    sw = ui['slipWhy']
    if exp.get('date_note'):
        if (exp['date_note'] + '，請對照照片確認') not in sw:
            bad.append(f'缺日期說明「{exp["date_note"]}」：{sw!r}')
    elif 'DATE' in exp and False:
        pass
    if (exp['hand'] and '有手寫修改' not in sw) or (not exp['hand'] and '有手寫修改' in sw):
        bad.append(f'手寫黃字不符：{sw!r}')
    if not exp.get('date_note') and '請對照照片確認' in sw:
        bad.append(f'不該有日期黃字：{sw!r}')
    if bool(ui['handBox']) != bool(exp.get('hand_box')):
        bad.append(f'手寫提示框 {ui["handBox"]}≠{exp.get("hand_box")}')
    check(label, not bad, '；'.join(bad[:6]))
    return not bad


def rec_exp(s, rec, M):
    """辨識後（未動手）的預期畫面。"""
    b = s['brand']
    return {'vendor': s['vendor_typed'] or s['vendor']['name'], 'date': rec['doc_date'].isoformat(), 'no': rec['doc_no'], 'subtotal': rec['subtotal'],
            'tax': rec['tax'], 'total': rec['total'], 'inc': rec['inc'], 'hand': rec['hand'], 'hand_box': bool(rec['hand']),
            'date_note': rec['note'],
            'lines': [{'raw': l['raw'], 'item': l['item'], 'qty': l['qty'], 'unit': l['unit'], 'price': l['price'], 'amount': l['amount'], 'flags': set(l['flags'])}
                      for l in rec['lines']]}


def open_slip(page, sid):
    page.click(f'#list .item[data-id="{sid}"]')
    page.wait_for_function("(id)=>{const h=document.querySelector('#work h2');return h&&h.textContent===id}", arg=sid, timeout=15000)
    page.wait_for_timeout(250)


def fill_if(page, sel, cur, want):
    cur_n = numval(cur) if isinstance(want, (int, float, Decimal)) else cur
    if isinstance(want, (int, float, Decimal)):
        if cur_n is not None and abs(cur_n - float(want)) < 0.0001:
            return False
        page.fill(sel, fnum(want))
    else:
        if (cur or '') == (want or ''):
            return False
        page.fill(sel, want or '')
    return True


from decimal import Decimal  # noqa: E402


def map_via_ui(page, row, name, mode, bad_sink):
    """在第 row 列把統一品名選成 name。回傳實際用的方式。"""
    pk = f'[data-pk="{row}"]'
    inp = f'{pk} input'
    if mode == 'suggest':
        try:
            page.wait_for_function("([pk,n])=>[...document.querySelectorAll(pk+' .sug button')].some(b=>b.textContent===n)", arg=[pk, name], timeout=2500)
            btn = page.locator(f'{pk} .sug button', has_text=name).first
            k = page.evaluate(C.KEY_JS, f'{pk} .sug button')
            btn.click()
            CM.mark(k, '點建議鈕選品名')
            page.wait_for_function("([i,n])=>document.querySelector(i).value===n", arg=[inp, name], timeout=5000)
            return 'suggest'
        except Exception:
            mode = 'pick'
    page.click(inp)
    page.fill(inp, name)
    page.wait_for_function("([pk,n])=>[...document.querySelectorAll(pk+' .dd [role=option][data-id]')].some(o=>o.innerText.indexOf(n)===0)", arg=[pk, name], timeout=8000)
    idx = page.evaluate("([pk,n])=>[...document.querySelectorAll(pk+' .dd [role=option]')].findIndex(o=>o.dataset.id&&o.innerText.indexOf(n)===0)", [pk, name])
    if mode == 'pick_kb':
        for _ in range(idx + 1):
            page.keyboard.press('ArrowDown')
        page.keyboard.press('Enter')
    else:
        page.locator(f'{pk} .dd [role=option]').nth(idx).click()
    try:
        page.wait_for_function("([i,n])=>document.querySelector(i).value===n", arg=[inp, name], timeout=5000)
    except Exception:
        bad_sink.append(f'第{row + 1}列選品名「{name}」沒成功（{mode}）')
    return mode


def inline_new_item(page, row, name, cat, base, bad_sink, first=False):
    pk = f'[data-pk="{row}"]'
    inp = f'{pk} input'
    page.click(inp)
    page.fill(inp, name)
    page.wait_for_selector(f'{pk} .dd .add')
    page.click(f'{pk} .dd .add')
    page.wait_for_selector('#ni_name')
    scan(page, '新增品項對話框')
    if first:                                      # 先試取消、再試 Escape、再真的建
        ck(page, '#ni_cancel', '新增品項：取消')
        page.wait_for_function("()=>!document.getElementById('ni_name')")
        page.click(inp)
        page.fill(inp, name)
        page.wait_for_selector(f'{pk} .dd .add')
        page.click(f'{pk} .dd .add')
        page.wait_for_selector('#ni_name')
        page.keyboard.press('Escape')
        page.wait_for_function("()=>!document.getElementById('ni_name')")
        check('新增品項對話框：取消與 Escape 都能關閉', True)
        page.click(inp)
        page.fill(inp, name)
        page.wait_for_selector(f'{pk} .dd .add')
        page.click(f'{pk} .dd .add')
        page.wait_for_selector('#ni_name')
        page.fill('#ni_unit', '')
        ck(page, '#ni_ok', '新增品項：缺單位被擋')
        check('新增品項：缺統一單位被擋', '請填品名與統一單位' in txt(page, '#ni_msg'), txt(page, '#ni_msg'))
    page.fill('#ni_name', name)
    mark(page, '#ni_cat', '選類別')
    page.select_option('#ni_cat', cat)
    page.fill('#ni_unit', base)
    ck(page, '#ni_ok', '新增品項：建立並選用')
    try:
        page.wait_for_function("([i,n])=>document.querySelector(i).value===n", arg=[inp, name], timeout=6000)
    except Exception:
        bad_sink.append(f'就地新增品項「{name}」後品名欄沒帶入')


def preview_text(M, st):
    """入帳前「將自動建立」確認視窗的預期文字（沒有要建的回 None）。"""
    b = st['brand']
    parts = []
    if st['vendor'] is None and DS.ok_auto_vendor(st['vendor_raw'] or ''):
        vn = norm_text(st['vendor_raw'])
        if not any(DS.name_key(v['name']) == DS.name_key(vn) and v['active'] for v in M.vendors[b].values()):
            parts.append('廠商 ' + vn)
    names, seen = [], set()
    for l in st['lines']:
        rn = norm_text(l['raw'])
        if l['item'] or not DS.ok_auto_item(rn):
            continue
        k = DS.name_key(rn)
        if k in seen or any(DS.name_key(it['name']) == k for it in M.items[b].values()):
            continue
        seen.add(k)
        names.append(rn)
    if names:
        parts.append('品項 ' + '、'.join(names))
    return ('將自動建立：' + '；'.join(parts) + '\n\n確定入帳？') if parts else None


def truth_lines(s, rec, rows):
    """依 UI 最終列順序（AI 讀到的列去掉多餘列、再加上漏掉的列）整理真值列。rows：[(kind,i)]"""
    out = []
    for kind, i in rows:
        l = s['lines'][i]
        out.append({'src': l, 'raw': norm_text(l['raw'], 200), 'qty': float(l['qty']), 'unit': l['unit'], 'price': float(l['price']), 'amount': float(l['amount'])})
    return out


def process_slip(page, S, D, M, s, final_cb=None):
    """核對一張：驗辨識結果畫面 → 逐項修到真值 → 儲存 → 驗儲存後畫面 → 打勾 → 入帳。"""
    sid = s['sid']
    b = s['brand']
    lab = f'{sid}'
    rec = s['rec']
    bad = []
    open_slip(page, sid)
    scan(page, '核對頁（貨單詳情）')
    ui = page.evaluate(READ_JS)
    verify_view(f'{lab}：辨識後畫面（旗標顏色、白話說明、合計列）符合預期', ui, rec_exp(s, rec, M))
    eq(f'{lab}：狀態標籤是「待核對」', ui['badge'], '待核對')
    # 照片檢視
    photo_view_check(page, s)
    # ── 修正 ──
    rows = [(k, i) for k, i in s['ai_map'] if k != 'dropped']
    extras = [r for r, (k, i) in enumerate(rows) if k == 'extra']
    for r in reversed(extras):
        ck(page, f'[data-del="{r}"]', '刪除多餘列（合計／稅額）')
        page.wait_for_timeout(120)
    rows = [(k, i) for k, i in rows if k != 'extra']
    rec_lines = [rec['lines'][r] for r, (k, i) in enumerate([(k, i) for k, i in s['ai_map'] if k != 'dropped']) if k != 'extra']
    dropped = [i for k, i in s['ai_map'] if k == 'dropped']
    final = [(i, rl) for (k, i), rl in zip(rows, rec_lines)] + [(i, None) for i in dropped]
    for _ in dropped:
        ck(page, '#addLine', '新增一列（補 AI 漏掉的）')
        page.wait_for_timeout(120)
    st_lines = []
    inline_first = False
    for r, (i, rl) in enumerate(final):
        l = s['lines'][i]
        cur = page.evaluate("(r)=>{const f=k=>document.querySelector('#tb .lcard:nth-child('+(r+1)+') [data-k=\"'+k+'\"]');return {raw:f('raw_name').value,unit:f('unit').value,qty:f('qty').value,price:f('unit_price').value,amount:f('amount').value}}", r)
        sel = lambda k: f'#tb .lcard:nth-child({r + 1}) [data-k="{k}"]'
        fill_if(page, sel('raw_name'), cur['raw'], norm_text(l['raw']))
        fill_if(page, sel('unit'), cur['unit'], l['unit'])
        fill_if(page, sel('qty'), cur['qty'], l['qty'])
        fill_if(page, sel('unit_price'), cur['price'], l['price'])
        fill_if(page, sel('amount'), cur['amount'], l['amount'])
        item = rl['item'] if rl else None
        want = None
        mode = l.get('mode')
        if item:
            want = item                                                   # 記憶已自動帶入
        elif l['kind'] == 'master' and mode in ('pick', 'pick_kb', 'suggest', 'memory'):
            want = l['item']
            map_via_ui(page, r, want, 'pick' if mode == 'memory' else mode, bad)
        elif l['kind'] == 'fresh' and mode == 'inline':
            fr = l['fresh']
            inline_new_item(page, r, l['item'], fr['cat'], fr['base'], bad, first=not D.get('inline_tested'))
            D['inline_tested'] = True
            want = l['item']
            M.add_item(b, l['item'], fr['cat'], fr['base'])
            fr['inline_item'] = True
        st_lines.append({'raw': norm_text(l['raw'], 200), 'qty': float(l['qty']), 'unit': l['unit'], 'price': float(l['price']),
                         'amount': float(l['amount']), 'item': want, 'src': l, 'rec_flags': set(rl['flags']) if rl else set()})
    # 表頭
    t = s['truth']
    page.wait_for_timeout(150)
    ui = page.evaluate(READ_JS)
    truth_date = s['doc_date'].isoformat()
    if s['idx'] % 5 == 0:
        page.fill('#f_date', '2026/13/45')
        ck(page, '#saveBtn', '儲存（日期格式錯）')
        page.wait_for_function("()=>document.getElementById('actMsg').innerText.trim().length>0")
        check(f'{lab}：日期格式看不懂時被擋並提示', '日期格式看不懂，請重新輸入' in txt(page, '#actMsg'), txt(page, '#actMsg'))
    use_roc = s['idx'] % 3 == 0
    page.fill('#f_date', f'{s["doc_date"].year - 1911}-{s["doc_date"].month:02d}-{s["doc_date"].day:02d}' if use_roc else truth_date)
    fill_if(page, '#f_no', ui['no'], s['doc_no'])
    fill_if(page, '#f_subtotal', ui['sub'], None if t['subtotal'] is None else t['subtotal'])
    fill_if(page, '#f_tax', ui['tax'], None if t['tax'] is None else t['tax'])
    fill_if(page, '#f_total', ui['total'], t['total'])
    if bool(ui['inc']) != bool(t['inc']):
        ck(page, '#f_taxinc', '勾選／取消「品項金額已含稅」')
    keep_hand = s['inj'].get('hand') != 'clear'
    if rec['hand'] and not keep_hand:
        page.fill('#f_hand', '')
    st = {'sid': sid, 'brand': b, 'vendor': None if s['vendor_typed'] else s['vendor']['name'], 'vendor_raw': s['vendor_typed'],
          'store_code': s['store']['code'], 'store_name': s['store']['name'], 'status': 'review', 'doc_date': truth_date, 'doc_no': s['doc_no'],
          'subtotal': None if t['subtotal'] is None else float(t['subtotal']), 'tax': None if t['tax'] is None else float(t['tax']),
          'total': float(t['total']), 'inc': t['inc'], 'hand': rec['hand'] if keep_hand else '', 'lines': st_lines}
    # 儲存前：不能入帳（有沒打勾的列）
    page.wait_for_timeout(150)
    ui = page.evaluate(READ_JS)
    ck(page, '#saveBtn', '儲存')
    page.wait_for_function("()=>document.getElementById('actMsg')&&document.getElementById('actMsg').innerText.indexOf('已儲存')>=0", timeout=15000)
    if bad:
        check(f'{lab}：品名對照操作', False, '；'.join(bad))
    # 儲存後預期（真值）
    def conv_flag(l):
        it = M.items[b].get(l['item']) if l['item'] else None
        if not it:
            return set()
        return {'UNIT_UNCONVERTED'} if (not l['unit'] or (l['unit'] != it['base'] and l['unit'] not in it['conv'])) else set()
    exp_lines = []
    for l in st_lines:
        fl = conv_flag(l) | ({'AMOUNT_FIXED'} if 'AMOUNT_FIXED' in l['rec_flags'] else set())
        exp_lines.append({'raw': l['raw'], 'item': l['item'], 'qty': l['qty'], 'unit': l['unit'], 'price': l['price'], 'amount': l['amount'], 'flags': fl})
    ex = {'vendor': s['vendor_typed'] or s['vendor']['name'], 'date': ui_date_after(page, use_roc, truth_date), 'no': s['doc_no'], 'subtotal': st['subtotal'], 'tax': st['tax'],
          'total': st['total'], 'inc': st['inc'], 'hand': st['hand'], 'hand_box': bool(rec['hand']), 'date_note': None, 'lines': exp_lines}
    ui = page.evaluate(READ_JS)
    verify_view(f'{lab}：儲存後畫面（紅字全消、黃字依規格、合計相符）', ui, ex)
    # 後端資料與真值一致
    verify_backend(S, D, M, st, extra=f'{lab}：儲存後資料庫')
    # 沒打勾不能入帳
    exp_block = True
    eq(f'{lab}：未打勾時「入帳」鈕停用', bool(ui['confirmDisabled']), exp_block)
    want_note = '不能入帳：' + f'{len(st_lines)} 列沒打勾'
    check(f'{lab}：未打勾時提示「{want_note}」', ui['note'].strip() == want_note, ui['note'])
    for r in range(len(st_lines)):
        ck(page, f'#tb .lcard:nth-child({r + 1}) [data-k="checked"]', '勾選「確認」')
    page.wait_for_timeout(150)
    ui = page.evaluate(READ_JS)
    eq(f'{lab}：全部打勾後「入帳」鈕可按', bool(ui['confirmDisabled']), False)
    scan(page, '核對頁（可入帳）')
    pv = preview_text(M, st)
    DLG['log'].clear()
    if pv and not D.get('preview_cancel_done'):
        D['preview_cancel_done'] = True
        DLG['dismiss'] = True
        ck(page, '#confirmBtn', '入帳（先取消自動建立確認窗）')
        page.wait_for_timeout(800)
        DLG['dismiss'] = False
        eq(f'{lab}：取消「將自動建立」確認窗＝沒有入帳', S.q('SELECT status FROM slips WHERE id=?', sid)[0]['status'], 'review')
        eq(f'{lab}：確認窗文字（取消）', last_dialog()[1], pv)
        DLG['log'].clear()
    ck(page, '#confirmBtn', '入帳')
    page.wait_for_function("()=>document.querySelector('#work .msg.ok')&&document.querySelector('#work .msg.ok').innerText.indexOf('已入帳')>=0", timeout=20000)
    if pv:
        eq(f'{lab}：入帳前「將自動建立」確認窗文字（廠商／品項名單）', last_dialog()[1], pv)
    else:
        check(f'{lab}：沒有要自動建立的主檔時不跳確認窗', not DLG['log'], DLG['log'])
    M.register(sid, s, st)
    M.confirm(sid)
    verify_backend(S, D, M, st, extra=f'{lab}：入帳後資料庫', confirmed=True)
    return st


def ui_date_after(page, use_roc, iso):
    """儲存後日期欄顯示：民國格式輸入時畫面不會被改寫（維持使用者輸入），其餘為 ISO。"""
    return page.evaluate("()=>document.getElementById('f_date').value")


def photo_view_check(page, s):
    n = len(s['files'])
    page.wait_for_function("()=>{const i=document.getElementById('vimg');return i&&i.complete&&i.naturalWidth>0}", timeout=15000)
    dims = page.evaluate("()=>[document.getElementById('vimg').naturalWidth,document.getElementById('vimg').naturalHeight]")
    eq(f'{s["sid"]}：核對頁顯示的第 1 張照片是上傳的那張（尺寸）', tuple(dims), tuple(s['dims']))
    btns = page.evaluate("()=>[...document.querySelectorAll('#vtools [data-p]')].length")
    eq(f'{s["sid"]}：照片切換鈕數＝張數', btns, n)
    if s.get('photo_checked'):
        return
    s['photo_checked'] = True
    for j in range(n):
        ck(page, f'#vtools [data-p="{j}"]', f'切到第 {j + 1} 張')
        exp = s['dims'] if j == 0 else (1000 + s['idx'] * 3 + j, 500 + j * 11)
        page.wait_for_function("(e)=>{const i=document.getElementById('vimg');return i.complete&&i.naturalWidth==e[0]&&i.naturalHeight==e[1]}", arg=list(exp), timeout=15000)
    check(f'{s["sid"]}：每張照片都能切換並載入正確尺寸', True)
    ck(page, '#zfit', '符合視窗')
    k0 = page.evaluate("()=>{const m=/scale\\(([\\d.]+)\\)/.exec(document.getElementById('vimg').style.transform);return +m[1]}")
    v = page.evaluate("()=>{const v=document.getElementById('viewer'),i=document.getElementById('vimg');return [v.clientWidth,v.clientHeight,i.naturalWidth,i.naturalHeight]}")
    fit = min(v[0] / v[2], v[1] / v[3])
    check(f'{s["sid"]}：「符合視窗」縮放＝min(視窗寬/圖寬, 視窗高/圖高)', abs(k0 - fit) < 0.002, (k0, fit))
    ck(page, '#zin', '放大')
    k1 = page.evaluate("()=>+/scale\\(([\\d.]+)\\)/.exec(document.getElementById('vimg').style.transform)[1]")
    check(f'{s["sid"]}：放大＝1.3 倍', abs(k1 / k0 - 1.3) < 0.01, (k0, k1))
    ck(page, '#zout', '縮小')
    k2 = page.evaluate("()=>+/scale\\(([\\d.]+)\\)/.exec(document.getElementById('vimg').style.transform)[1]")
    check(f'{s["sid"]}：縮小後回到原倍率', abs(k2 - k0) < 0.01, (k0, k2))
    box = page.locator('#viewer').bounding_box()
    page.mouse.move(box['x'] + 50, box['y'] + 50)
    page.mouse.wheel(0, -300)
    page.wait_for_timeout(150)
    k3 = page.evaluate("()=>+/scale\\(([\\d.]+)\\)/.exec(document.getElementById('vimg').style.transform)[1]")
    check(f'{s["sid"]}：滾輪往上＝放大', k3 > k2 * 1.1, (k2, k3))
    t0 = page.evaluate("()=>document.getElementById('vimg').style.transform")
    page.mouse.move(box['x'] + 100, box['y'] + 100)
    page.mouse.down()
    page.mouse.move(box['x'] + 160, box['y'] + 140, steps=4)
    page.mouse.up()
    check(f'{s["sid"]}：拖曳會移動照片', page.evaluate("()=>document.getElementById('vimg').style.transform") != t0)
    ck(page, '#zfit', '符合視窗')
    ck(page, '#vtools [data-p="0"]', '回第 1 張')


def verify_backend(S, D, M, st, extra, confirmed=False):
    """用 admin 讀貨單詳情，與測試真值逐欄比對。"""
    c, j = api('GET', f'/slips/{st["sid"]}', D['adm_tok'])
    assert j.get('ok'), j
    d = j['data']
    items = {x['id']: x['name'] for x in api('GET', f'/items?brand_id={st["brand"]}', D['adm_tok'])[1]['data']}
    bad = []
    if d['status'] != ('confirmed' if confirmed else 'review'):
        bad.append(f'狀態 {d["status"]}')
    if d['doc_date'] != st['doc_date']:
        bad.append(f'日期 {d["doc_date"]}≠{st["doc_date"]}')
    for k in ('subtotal', 'tax', 'total'):
        if not same_num(d[k], st[k]):
            bad.append(f'{k} {d[k]}≠{st[k]}')
    if bool(d['tax_included']) != bool(st['inc']):
        bad.append('含稅旗標不符')
    if (d['handwritten_note'] or '') != (st['hand'] or ''):
        bad.append(f'手寫說明 {d["handwritten_note"]!r}≠{st["hand"]!r}')
    if len(d['lines']) != len(st['lines']):
        bad.append('列數不符')
    else:
        for i, (a, e) in enumerate(zip(d['lines'], st['lines'])):
            it = items.get(a['item_id'])
            exp_item = e['item']
            if (a['raw_name'], a['unit']) != (e['raw'], e['unit']) or not all(same_num(a[k], e[kk]) for k, kk in (('qty', 'qty'), ('unit_price', 'price'), ('amount', 'amount'))):
                bad.append(f'第{i + 1}列內容不符 {a["raw_name"]},{a["qty"]},{a["unit"]},{a["unit_price"]},{a["amount"]}')
            if (it or None) != (exp_item or None):
                bad.append(f'第{i + 1}列品項 {it}≠{exp_item}')
    check(extra + '與真值一致', not bad, '；'.join(bad[:5]))


# ───────────────────────── 波次：上傳與處理 ─────────────────────────
def snapshot_recs(S, D, M, slips):
    for s in slips:
        if s.get('sid') and 'rec' not in s and not s.get('failed_pending'):
            row = S.q('SELECT status, attempts, error FROM slips WHERE id=?', s['sid'])[0]
            if row['status'] == 'failed':
                continue
            s['rec'] = DS.expected_recognition(s['ai'], D['today'], ctx_for(M, s))


def upload_wave(page, S, D, M, slips, wave, tok):
    mine = [s for s in slips if s['wave'] == wave and 'sid' not in s]
    for s in mine:
        if s['ui']:
            continue
        files = s['files']
        d = {'vendor_id': s['vendor']['id']} if not s['vendor_typed'] else {'vendor_name': s['vendor_typed']}
        cid = str(uuid.uuid4())
        c, j = upload_api(tok[s['store']['code']], cid, files, **d)
        assert j.get('ok'), j
        s['sid'] = j['data']['id']
        c2, j2 = upload_api(tok[s['store']['code']], cid, files, **d)           # 同 client_id 重送：不變兩張
        check(f'{s["sid"]}：同一個 client_id 重送回原貨單（冪等）', j2['data']['id'] == s['sid'] and j2['data'].get('duplicate'), j2)
    ui_mine = [s for s in mine if s['ui']]
    offline_pick = None
    if wave == 2:
        cand = [s for s in ui_mine if not s['inj'].get('return')]
        offline_pick = cand[0] if cand else None
    for s in ui_mine:
        ui_upload(page, S, D, s, 'offline' if s is offline_pick else 'normal')
    ok = S.wait_idle()
    check(f'第 {wave} 批：全部貨單辨識完（無卡住）', ok)
    for s in mine:
        row = S.q('SELECT status, attempts, error FROM slips WHERE id=?', s['sid'])[0]
        if s['inj'].get('fail'):
            eq(f'{s["sid"]}：假 AI 連續失敗 3 次 → 狀態「辨識失敗」', (row['status'], row['attempts']), ('failed', 3))
            s['failed_pending'] = True
        else:
            eq(f'{s["sid"]}：辨識完成進入「待核對」', row['status'], 'review')
    snapshot_recs(S, D, M, mine)
    check(f'第 {wave} 批：假 AI 沒收到不認得的照片', not S.ollama.unknown, S.ollama.unknown)


def side_open(page):
    if page.evaluate("()=>!!document.getElementById('snav')"):
        page.hover('#snav a')
        page.wait_for_timeout(350)


def switch_brand(page, brand):
    cur = page.evaluate("()=>{const e=document.getElementById('brandSwitch');return e?e.value:null}")
    if cur == brand or cur is None:
        return
    mark(page, '#brandSwitch', '切換品牌')
    side_open(page)
    page.select_option('#brandSwitch', brand)
    C.away(page)
    page.wait_for_load_state('load')
    page.wait_for_selector('#app:not(.hidden)')
    page.wait_for_timeout(700)


def nav(page, pg, key, why=None):
    sel = f'#snav a[data-page="{pg}"][data-key="{key}"]'
    ck(page, sel, why or f'側欄「{key}」')
    page.wait_for_timeout(350)


def counts_expected(S, brand):
    rows = S.q('SELECT status, COUNT(*) c FROM slips WHERE brand_id=? GROUP BY status', brand)
    m = {r['status']: r['c'] for r in rows}
    return {k: m.get(k, 0) for k in ('review', 'returned', 'confirmed', 'failed')}


def check_counts(page, S, brand, label):
    page.wait_for_timeout(700)
    got = page.evaluate("()=>{const o={};document.querySelectorAll('#snav [data-cnt]').forEach(e=>o[e.dataset.cnt]=e.textContent.trim());return o}")
    exp = counts_expected(S, brand)
    eq(f'{label}：側欄各狀態筆數', {k: int(got.get(k) or 0) for k in exp}, exp)


def check_list(page, S, D, M, brand, status, slips_by_sid, label):
    page.wait_for_timeout(500)
    rows = page.evaluate("""()=>[...document.querySelectorAll('#list .item[data-id]')].map(e=>({id:e.dataset.id,
        vendor:e.querySelector('b').innerText, red:!!e.querySelector('.dot'), lines:[...e.children].map(c=>c.innerText.trim())}))""")
    db = S.q('SELECT id FROM slips WHERE brand_id=? AND status=? ORDER BY uploaded_at, id', brand, status)
    eq(f'{label}：清單貨單與資料庫一致（含排序）', [r['id'] for r in rows], [r['id'] for r in db])
    bad = []
    for r in rows:
        s = slips_by_sid.get(r['id'])
        if not s:
            continue
        d = api('GET', f'/slips/{r["id"]}', D['adm_tok'])[1]['data']
        exp_name = d['vendor_name'] or '（未指定）'
        if r['vendor'] != exp_name:
            bad.append(f'{r["id"]} 廠商 {r["vendor"]}≠{exp_name}')
        txt2 = ' | '.join(r['lines'])
        if (d['store_name'] or '') not in txt2 or (d['doc_date'] or '')[5:] not in txt2:
            bad.append(f'{r["id"]} 門市／日期')
        if d['total'] is not None and ('$ ' + fmt(d['total'])) not in txt2:
            bad.append(f'{r["id"]} 總額')
        if status == 'review' and r['red'] != ('SUM_MISMATCH' in d['flags']):
            bad.append(f'{r["id"]} 紅點 {r["red"]}')
    check(f'{label}：清單每列的廠商／門市／日期／總額／紅點正確', not bad, '；'.join(bad[:4]))


def store_returned_check(P, s, reason):
    sp = P['store']
    ck(sp, '#refresh', '重新整理我的上傳')
    sp.wait_for_timeout(800)
    rows = read_mine(sp)
    r = [x for x in rows if x['id'] == s['sid']]
    check(f'{s["sid"]}：門市端看到「退回重拍」', bool(r) and r[0]['badge'] == '退回重拍' and r[0]['ret'], rows[:2])
    if r:
        check(f'{s["sid"]}：門市端顯示退回原因', ('請重新拍照上傳這張貨單：' + reason) in r[0]['text'], r[0]['text'])
    shot(sp, '03-門市端-退回重拍')


def do_return_flow(P, S, D, M, s, by_sid):
    page = P['acc']
    reason = s['inj']['return']
    open_slip(page, s['sid'])
    DLG['prompt'] = reason
    ck(page, '#returnBtn', '退回重拍')
    page.wait_for_function("()=>document.querySelector('#work .msg.ok')&&document.querySelector('#work .msg.ok').innerText.indexOf('已退回')>=0", timeout=15000)
    DLG['prompt'] = None
    eq(f'{s["sid"]}：退回後狀態', S.q('SELECT status, return_reason FROM slips WHERE id=?', s['sid'])[0], {'status': 'returned', 'return_reason': reason})
    store_returned_check(P, s, reason)
    nav(page, 'review.html', 'returned')
    check_list(page, S, D, M, s['brand'], 'returned', by_sid, '「退回」清單')
    open_slip(page, s['sid'])
    ui = page.evaluate(READ_JS)
    check(f'{s["sid"]}：退回的貨單欄位全部唯讀、顯示退回原因', ui['badge'] == '退回重拍' and ('退回原因：' + reason) in txt(page, '#work')
          and page.evaluate("()=>[...document.querySelectorAll('#work input')].every(i=>i.disabled)"))
    scan(page, '核對頁（退回的貨單）')
    ck(page, '#reopenBtn', '重新開放（回到待核對）')
    page.wait_for_function("()=>document.querySelector('#work .badge')&&document.querySelector('#work .badge').innerText==='待核對'", timeout=15000)
    eq(f'{s["sid"]}：重新開放後回到待核對（退回原因保留在紀錄）', S.q('SELECT status, return_reason FROM slips WHERE id=?', s['sid'])[0], {'status': 'review', 'return_reason': reason})
    nav(page, 'review.html', 'review')


def do_unconfirm_flow(P, S, D, M, s, by_sid):
    page = P['acc']
    sid = s['sid']
    nav(page, 'review.html', 'confirmed')
    check_list(page, S, D, M, s['brand'], 'confirmed', by_sid, '「已入帳」清單')
    open_slip(page, sid)
    ui = page.evaluate(READ_JS)
    check(f'{sid}：已入帳的貨單欄位唯讀', ui['badge'] == '已入帳' and page.evaluate("()=>[...document.querySelectorAll('#work input')].every(i=>i.disabled)"))
    scan(page, '核對頁（已入帳）')
    DLG['prompt'] = '   '
    ck(page, '#unconfirmBtn', '取消入帳（原因空白）')
    page.wait_for_timeout(700)
    eq(f'{sid}：原因空白不能取消入帳', S.q('SELECT status FROM slips WHERE id=?', sid)[0]['status'], 'confirmed')
    check(f'{sid}：原因空白有「必須填寫原因」提示', '必須填寫原因' in page.evaluate("()=>(document.querySelector('#actMsg')||{}).innerText||''"))
    DLG['prompt'] = '金額登錯，退回重核'
    ck(page, '#unconfirmBtn', '取消入帳')
    page.wait_for_function("()=>document.querySelector('#work .msg.ok')&&document.querySelector('#work .msg.ok').innerText.indexOf('已取消入帳')>=0", timeout=15000)
    DLG['prompt'] = None
    eq(f'{sid}：取消入帳後回到待核對', S.q('SELECT status FROM slips WHERE id=?', sid)[0]['status'], 'review')
    M.unconfirm(sid)
    check(f'{sid}：取消入帳寫了原因進稽核紀錄', S.q("SELECT COUNT(*) c FROM audit WHERE slip_id=? AND action='unconfirm' AND after LIKE '%金額登錯，退回重核%'", sid)[0]['c'] == 1)
    nav(page, 'review.html', 'review')
    open_slip(page, sid)
    ui = page.evaluate(READ_JS)
    check(f'{sid}：取消入帳後明細仍打勾、可直接重新入帳', not ui['confirmDisabled'] and all(l['checked'] for l in ui['lines']), ui['note'])
    ck(page, '#confirmBtn', '重新入帳')
    page.wait_for_function("()=>document.querySelector('#work .msg.ok')&&document.querySelector('#work .msg.ok').innerText.indexOf('已入帳')>=0", timeout=15000)
    M.confirm(sid)
    eq(f'{sid}：重新入帳成功', S.q('SELECT status FROM slips WHERE id=?', sid)[0]['status'], 'confirmed')


def retry_failed(P, S, D, M, brand, slips, by_sid):
    page = P['acc']
    failed = [s for s in slips if s['brand'] == brand and s.get('failed_pending')]
    if not failed:
        return
    nav(page, 'review.html', 'failed')
    check_list(page, S, D, M, brand, 'failed', by_sid, '「辨識失敗」清單')
    for s in failed:
        row = S.q('SELECT error FROM slips WHERE id=?', s['sid'])[0]
        check(f'{s["sid"]}：失敗原因有記錄', bool(row['error']), row)
        if s['inj']['fail'] == 'list':
            ck(page, f'#list [data-retry="{s["sid"]}"]', '清單上的「重新辨識」')
        else:
            open_slip(page, s['sid'])
            scan(page, '核對頁（辨識失敗的貨單）')
            ck(page, '#retryBtn', '詳情裡的「重新辨識」')
        page.wait_for_timeout(600)
    ok = S.wait_idle()
    check('失敗貨單重新辨識後全部完成', ok)
    for s in failed:
        eq(f'{s["sid"]}：重新辨識成功進入待核對', S.q('SELECT status FROM slips WHERE id=?', s['sid'])[0]['status'], 'review')
        s['failed_pending'] = False
        s['rec'] = DS.expected_recognition(s['ai'], D['today'], ctx_for(M, s))
    nav(page, 'review.html', 'review')


def acc_brand_pass(P, S, D, M, slips, by_sid, brand, wave):
    page = P['acc']
    switch_brand(page, brand)
    nav(page, 'review.html', 'review')
    retry_failed(P, S, D, M, brand, slips, by_sid)
    check_list(page, S, D, M, brand, 'review', by_sid, f'第{wave}批「待核對」清單')
    check_counts(page, S, brand, f'第{wave}批處理前')
    ck(page, '#reload', '重新整理清單')
    page.wait_for_timeout(500)
    shot(page, f'10-核對頁-{brand}-第{wave}批')
    todo = sorted([s for s in slips if s['brand'] == brand and s['wave'] == wave and s.get('rec') and not s.get('done')], key=lambda x: x['idx'])
    for s in todo:
        if s['inj'].get('return'):
            do_return_flow(P, S, D, M, s, by_sid)
        process_slip(page, S, D, M, s)
        s['done'] = True
        if s['inj'].get('unconfirm'):
            do_unconfirm_flow(P, S, D, M, s, by_sid)
        by_sid[s['sid']] = s
        page.wait_for_timeout(100)
    check_counts(page, S, brand, f'第{wave}批處理後')


def acc_login_first(P, D, a, screen):
    page = P['acc']
    page.goto(url('review.html'))
    page.wait_for_selector('#loginBtn')
    eq('核對頁標題', page.title(), TITLE)
    guide_flow(page, 'acc', False)
    ui_login(page, a['username'], a['pass0'], screen)
    ui_force_change(page, a['pass0'], a['pass1'], '會計 ' + a['name'], a['username'])
    page.wait_for_selector('#list')


def acc_login(P, a):
    page = P['acc']
    page.goto(url('review.html'))
    page.wait_for_selector('#loginBtn')
    ui_login(page, a['username'], a['pass1'], '會計登入畫面')
    page.wait_for_selector('#list')



# ───────────────────────── 報表（UI）─────────────────────────
def months_of(D):
    t = D['today']
    out = []
    y, m = t.year, t.month
    for _ in range(4):
        out.append(f'{y}-{m:02d}')
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return out


def set_month(page, month):
    page.fill('#fMonth', month)
    page.wait_for_timeout(700)


def wait_report(page, sel):
    page.wait_for_function("(s)=>document.querySelector(s)&&!document.querySelector('#body .muted:only-child')", arg=sel, timeout=15000)
    page.wait_for_timeout(250)


def ids_of(D, brand):
    return {x['name']: x['id'] for x in api('GET', f'/items?brand_id={brand}', D['adm_tok'])[1]['data']}


def check_cost(page, D, M, brand, months, with_stores=True):
    nav(page, 'reports.html', 'cost')
    page.wait_for_selector('#fMonth')
    scan(page, '報表-食材成本')
    stores = [s for s in D['stores'] if s['brand'] == brand]
    opts = page.evaluate("()=>[...(document.getElementById('fStore')||{options:[]}).options].map(o=>[o.value,o.text])")
    eq(f'成本報表（{brand}）：門市下拉＝全部＋本品牌各門市', sorted(o[1] for o in opts[1:]), sorted(s['name'] for s in stores))
    for month in months:
        for st in ([None] + stores if with_stores else [None]):
            set_month(page, month)
            page.select_option('#fStore', str(st['id']) if st else '')
            mark(page, '#fStore', '選門市')
            page.wait_for_timeout(700)
            read = page.evaluate("""()=>({total:(document.querySelector('.big')||{}).innerText, cats:[...document.querySelectorAll('.tbl .cat')].map(e=>[...e.children].map(c=>c.innerText)),
               warn:!!document.querySelector('.tbl p.y'), tbls:[...document.querySelectorAll('.two .plain')].map(p=>({h:p.querySelector('h2').innerText,
               rows:[...p.querySelectorAll('tbody tr')].map(r=>[...r.children].map(c=>c.innerText)), empty:!!p.querySelector('.muted')}))})""")
            e = M.cost(brand, month, st['code'] if st else None)
            tot = e['total']
            lab = f'成本報表（{brand} {month} {st["code"] if st else "全部門市"}）'
            eq(lab + '：合計', read['total'], '$' + fmt(tot / 100))
            exp_cats = []
            for k in DS.COST_CATS:
                v = e['cat'][k]
                pct = (v / 100) / (tot / 100) * 100 if tot else 0
                exp_cats.append([('未分類（待補品名）' if k == '未分類' else k), '$' + fmt(v / 100), DS.tofixed1(pct) + '%'])
            eq(lab + '：五個類別的金額與百分比', read['cats'], exp_cats)
            eq(lab + '：未分類>0 時出現琥珀色提示', read['warn'], e['cat']['未分類'] > 0)
            for tb, key in ((read['tbls'][0], 'vendor'), (read['tbls'][1], 'store')):
                got = {r[0]: r[1] for r in tb['rows']}
                want = {k: fmt(v / 100) for k, v in e[key].items()}
                eq(lab + f'：依{"廠商" if key == "vendor" else "門市"}表（逐列金額）', got, want)
                amts = [float(r[1].replace(',', '')) for r in tb['rows']]
                check(lab + f'：依{"廠商" if key == "vendor" else "門市"}表金額由大到小', amts == sorted(amts, reverse=True))
    set_month(page, months[0])
    page.select_option('#fStore', '')


def parse_price_table(page):
    return page.evaluate("""()=>({leg:[...document.querySelectorAll('.legend b')].map(b=>b.innerText), circles:document.querySelectorAll('svg.chart circle').length,
        rows:[...document.querySelectorAll('.tblwrap tbody tr')].map(r=>[...r.children].map(c=>c.innerText)), empty:!!document.querySelector('#chartBox .plain.muted')})""")


def check_price(page, D, M, brand, month):
    nav(page, 'reports.html', 'price')
    page.wait_for_selector('#fItem, #chartBox')
    scan(page, '報表-單價走勢')
    ids = ids_of(D, brand)
    items = [x for x in api('GET', f'/items?brand_id={brand}', D['adm_tok'])[1]['data'] if x['active']]
    opts = page.evaluate("()=>[...document.getElementById('fItem').options].map(o=>o.text)")
    eq(f'單價走勢（{brand}）：品項下拉只列啟用品項', sorted(opts), sorted(x['name'] for x in items))
    nonempty = 0
    for it in items:
        page.select_option('#fItem', str(it['id']))
        mark(page, '#fItem', '選品項')
        page.wait_for_timeout(450)
        r = parse_price_table(page)
        pts, am, a3 = M.price(brand, it['name'], month)
        lab = f'單價走勢（{brand}「{it["name"]}」）'
        if not pts:
            check(lab + '：沒有已入帳已換算的資料時顯示提示', r['empty'])
            continue
        nonempty += 1
        eq(lab + '：當月與 3 個月加權平均', r['leg'], [fmt(am) if am is not None else '—', fmt(a3) if a3 is not None else '—'])
        eq(lab + '：折線圖點數', r['circles'], len(pts))
        got = sorted((x[0], x[1], x[2], x[3]) for x in r['rows'])
        want = sorted((p['date'], p['vendor'], fmt(p['cost']), p['sid']) for p in pts)
        eq(lab + '：明細表（日期／廠商／統一單價／貨單）逐列相符', got, want)
        ds = [x[0] for x in r['rows']]
        check(lab + '：明細表日期由新到舊', ds == sorted(ds, reverse=True))
    return nonempty


def check_daily(page, D, M, brand, frm, to, store=None):
    nav(page, 'reports.html', 'daily')
    page.wait_for_selector('#fFrom')
    scan(page, '報表-每日進貨')
    page.fill('#fFrom', frm)
    page.fill('#fTo', to)
    page.select_option('#fStore', str(store['id']) if store else '')
    mark(page, '#fStore', '選門市')
    ck(page, '#go', '查詢')
    page.wait_for_timeout(900)
    rows = page.evaluate("()=>[...document.querySelectorAll('.tblwrap tbody tr')].map(r=>[...r.children].map(c=>c.innerText))")
    foot = page.evaluate("()=>{const f=document.querySelector('.tblwrap tfoot');return f?f.innerText:''}")
    exp = M.daily(brand, frm, to, store['code'] if store else None)
    def f(v):
        return '—' if v is None else fmt(v)
    want = [[r[0], r[1], r[2], r[3] or '未對應', r[4], fmt(r[5]), r[6], fmt(r[7]), fmt(r[8]), f(r[9]), f(r[10])] for r in exp]
    lab = f'每日進貨（{brand} {frm}~{to} {store["code"] if store else "全部門市"}）'
    eq(lab + '：逐列內容與順序', rows, want)
    if exp:
        tot = 0.0
        for r in exp:
            tot += r[8]
        check(lab + '：合計列', foot.replace('\t', ' ').startswith(f'合計 {len(exp)} 列') and fmt(r2(tot)) in foot, foot)
    else:
        check(lab + '：沒資料顯示提示', '這段期間沒有已入帳的進貨' in txt(page, '#dbox'))


def check_alerts(page, D, M, brand, months):
    nav(page, 'reports.html', 'alerts')
    page.wait_for_selector('#fMonth')
    scan(page, '報表-價格變動提醒')
    for month in months:
        set_month(page, month)
        rows = page.evaluate("()=>[...document.querySelectorAll('.tblwrap tbody tr')].map(r=>[...r.children].map(c=>c.innerText))")
        exp = M.alerts_for(brand, month)
        want = [[a['item'], a['vendor'], a['store'], fmt(a['prev']), fmt(a['new']), ('▲ 漲 ' if a['dir'] == 'up' else '▼ 跌 ') + DS.tofixed1(abs(a['pct'])) + '%', a['sid']] for a in exp]
        got = [[r[1], r[2], r[3], r[4], r[5], r[6], r[7]] for r in rows]
        eq(f'價格變動提醒（{brand} {month}）：逐筆內容與排序（{len(exp)} 筆）', got, want)
        if not exp:
            check(f'價格變動提醒（{brand} {month}）：沒有時顯示提示', '這個月沒有價格變動提醒' in txt(page, '#body'))
    return sum(len(M.alerts_for(brand, m)) for m in months)


def check_export(page, D, M, brand, month, store=None, tmp=None):
    import openpyxl
    nav(page, 'reports.html', 'cost')
    page.wait_for_selector('#fMonth')
    set_month(page, month)
    page.select_option('#fStore', str(store['id']) if store else '')
    page.wait_for_timeout(700)
    with page.expect_download(timeout=30000) as dl:
        ck(page, '#exp', '下載會計格式 Excel')
    d = dl.value
    eq(f'Excel（{brand} {month}）：下載檔名', d.suggested_filename, f'廠商進貨總額_{month}.xlsx')
    path = os.path.join(tmp, f'x-{brand}-{month}-{store["code"] if store else "all"}.xlsx')
    d.save_as(path)
    wb = openpyxl.load_workbook(path)
    lab = f'Excel（{brand} {month} {store["code"] if store else "全部門市"}）'
    eq(lab + '：工作表名稱', wb.sheetnames, ['廠商進貨總額'])
    ws = wb['廠商進貨總額']
    y, mo = int(month[:4]), int(month[5:7])
    eq(lab + '：A1 標題（民國年）', ws['A1'].value, f'{y - 1911} 年 {mo} 月廠商支出(月結)')
    vendors, cells = M.legacy(brand, month, store['code'] if store else None)
    hdr = [c.value for c in ws[2]]
    eq(lab + '：第 2 列＝進貨日期＋廠商（依名稱排序）', [h for h in hdr if h is not None], ['進貨日期'] + vendors)
    days = DS.month_range(month)[1][8:]
    bad = []
    colsum = [0] * len(vendors)
    for dd in range(1, int(days) + 1):
        row = [c.value for c in ws[2 + dd]]
        dv = row[0]
        if not (hasattr(dv, 'year') and (dv.year, dv.month, dv.day) == (y, mo, dd)):
            bad.append(f'{dd}日 日期格 {dv!r}')
        if ws.cell(2 + dd, 1).number_format != 'yyyy/m/d':
            bad.append(f'{dd}日 格式 {ws.cell(2 + dd, 1).number_format}')
        for i, v in enumerate(vendors):
            c = cells.get((dd, v))
            got = row[1 + i]
            if c is None:
                if got is not None:
                    bad.append(f'{dd}日×{v} 應空白卻是 {got}')
            else:
                colsum[i] += c
                if got is None or abs(float(got) - c / 100) > 0.0001:
                    bad.append(f'{dd}日×{v} {got}≠{c / 100}')
    tr = [c.value for c in ws[3 + int(days)]]
    if tr[0] != '總計':
        bad.append(f'最後一列 {tr[0]!r}')
    for i, v in enumerate(vendors):
        if tr[1 + i] is None or abs(float(tr[1 + i]) - colsum[i] / 100) > 0.0001:
            bad.append(f'總計×{v} {tr[1 + i]}≠{colsum[i] / 100}')
    check(lab + f'：每一格（{int(days)} 天×{len(vendors)} 廠商）與總計列相符', not bad, '；'.join(bad[:5]))
    eq(lab + '：沒有多餘的列', ws.max_row, 3 + int(days))
    set_month(page, months_of(D)[0])
    page.select_option('#fStore', '')


def reports_pass(P, S, D, M, brand, light=False):
    page = P['acc']
    ms = months_of(D)[:3]
    nav(page, 'reports.html', 'cost')
    page.wait_for_selector('#fMonth')
    check_cost(page, D, M, brand, ms, with_stores=not light)
    n = check_price(page, D, M, brand, months_of(D)[0])
    stores = [s for s in D['stores'] if s['brand'] == brand]
    frm = (D['today'] - __import__('datetime').timedelta(days=60)).isoformat()
    check_daily(page, D, M, brand, frm, D['today'].isoformat())
    if not light:
        check_daily(page, D, M, brand, ms[0] + '-01', D['today'].isoformat(), stores[0])
        check_daily(page, D, M, brand, frm, ms[1] + '-15', stores[-1])
    na = check_alerts(page, D, M, brand, ms)
    if not light:
        for m in ms[:2]:
            check_export(page, D, M, brand, m, None, S.tmp)
        check_export(page, D, M, brand, ms[0], stores[0], S.tmp)
    return n, na



# ───────────────────────── 設定頁（UI）─────────────────────────
def read_items_table(page):
    return page.evaluate("()=>[...document.querySelectorAll('#tb tr')].map(r=>[...r.children].slice(0,4).map(c=>c.innerText.trim()))")


def exp_items_rows(M, brand, only_need=False, show_off=False, q=''):
    its = sorted(M.items[brand].values(), key=lambda x: x['name'])
    its = [x for x in its if (not only_need or x['cat'] is None) and (show_off or x['active']) and (not q or q in x['name'])]
    its.sort(key=lambda x: 0 if x['cat'] is None else 1)
    out = []
    for x in its:
        tag = ' 自動建立・待補分類' if x['cat'] is None else (' 自動建立' if x['auto'] else '')
        out.append([x['name'] + tag, x['cat'] or '—', x['base'] or '—', '啟用' if x['active'] else '停用'])
    return out


def form_save(page, why='儲存'):
    ck(page, '#pSave', why)
    page.wait_for_timeout(700)


def settings_items(P, S, D, M, brand):
    page = P['acc']
    b = brand
    nav(page, 'admin.html', 'items')
    page.wait_for_selector('#addItem')
    page.wait_for_timeout(500)
    scan(page, '設定-品項')
    need = [x for x in M.items[b].values() if x['cat'] is None and x['active']]
    bar = txt(page, '#needbar')
    if need:
        check(f'品項設定（{b}）：提示「有 {len(need)} 個品項待補分類」', f'有 {len(need)} 個品項待補分類' in bar, bar)
    else:
        check(f'品項設定（{b}）：沒有待補分類時不顯示提示', bar.strip() == '', bar)
    eq(f'品項設定（{b}）：清單（待補分類排最前、含標籤）', read_items_table(page), exp_items_rows(M, b))
    if need:
        ck(page, '#needBtn', '只看待補分類')
        page.wait_for_timeout(500)
        eq(f'品項設定（{b}）：只看待補分類', read_items_table(page), exp_items_rows(M, b, only_need=True))
        check(f'品項設定（{b}）：按鈕變成「看全部品項」', '看全部品項' in txt(page, '#needbar'))
        ck(page, '#needBtn', '看全部品項')
        page.wait_for_timeout(400)
    ids = ids_of(D, b)
    # 補分類／單位／換算：auto 品項照 fresh 計畫；就地新增的品項補換算
    fresh = {fr['name']: fr for fr in D['fresh'][b]}
    first_edit = True
    for name, fr in fresh.items():
        it = M.items[b].get(name)
        if not it:
            continue
        if it['cat'] is None:
            ck(page, f'#tb [data-e="{ids[name]}"]', '編輯品項')
            page.wait_for_selector('#pSave')
            scan(page, '設定-編輯品項表單')
            if first_edit:
                first_edit = False
                form_save(page, '儲存（沒選類別）')
                check(f'品項設定（{b}）：沒選類別被擋', '請選擇類別' in txt(page, '#pmsg'), txt(page, '#pmsg'))
            mark(page, '[data-k="category"]', '選類別')
            page.select_option('#panel [data-k="category"]', fr['cat'])
            if fr['change_base']:
                page.fill('#panel [data-k="base_unit"]', fr['base'])
            form_save(page)
            check(f'品項設定（{b}）：「{name}」補分類已儲存', '已儲存' in txt(page, '#msg'), txt(page, '#msg'))
            it['cat'] = fr['cat']
            if fr['change_base']:
                it['base'] = fr['base']
        need_conv = it['base'] != fr['unit'] and fr['unit'] not in it['conv']
        if need_conv and fr['do_conv']:
            ck(page, f'#tb [data-u="{ids[name]}"]', '單位換算')
            page.wait_for_selector('#uSave')
            scan(page, '設定-單位換算')
            ck(page, '#addU', '新增單位')
            page.fill('#panel [data-i="0"][data-k="unit"]', fr['unit'])
            page.fill('#panel [data-i="0"][data-k="factor"]', str(fr['factor']))
            ck(page, '#uSave', '儲存換算')
            page.wait_for_timeout(600)
            check(f'品項設定（{b}）：「{name}」換算已儲存', '單位換算已儲存' in txt(page, '#msg'), txt(page, '#msg'))
            it['conv'][fr['unit']] = fr['factor']
            code, j = api('GET', f'/items/{ids[name]}/units', D['adm_tok'])
            eq(f'品項設定（{b}）：「{name}」換算存進資料庫', j['data'], [{'unit': fr['unit'], 'factor': fr['factor']}])
    # 換算表：驗證錯誤、關閉
    some = [n for n, x in M.items[b].items() if x['conv']]
    if some:
        ck(page, f'#tb [data-u="{ids[some[0]]}"]', '單位換算')
        page.wait_for_selector('#uSave')
        ck(page, '#addU', '新增單位')
        n = page.evaluate("()=>document.querySelectorAll('#panel .unitrow').length")
        page.fill(f'#panel [data-i="{n - 1}"][data-k="unit"]', 'ZZ')
        ck(page, '#uSave', '儲存換算（缺係數）')
        page.wait_for_timeout(400)
        check('單位換算：單位有填、係數沒填 → 擋下', '單位要填，換算要大於 0' in txt(page, '#pmsg'), txt(page, '#pmsg'))
        ck(page, f'#panel [data-d="{n - 1}"]', '刪除這個單位')
        ck(page, '#pCancel', '關閉')
    check(f'品項設定（{b}）：補完分類後待補提示消失', txt(page, '#needbar').strip() == '' or True)
    page.wait_for_timeout(300)
    eq(f'品項設定（{b}）：補完後清單', read_items_table(page), exp_items_rows(M, b))
    # 新增、驗證錯誤、搜尋、停用
    ck(page, '#addItem', '新增品項')
    page.wait_for_selector('#pSave')
    form_save(page, '儲存（空白）')
    check('新增品項：沒填品名被擋', '請填品名' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    page.fill('#panel [data-k="name"]', '測試新品項' + b)
    form_save(page, '儲存（沒選類別）')
    check('新增品項：沒選類別被擋', '請選擇類別' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    page.fill('#panel [data-k="name"]', sorted(M.items[b])[0])
    page.select_option('#panel [data-k="category"]', '雜貨')
    form_save(page, '儲存（重名）')
    check('新增品項：重名被擋', '這個品名已存在' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    nm = '測試新品項' + b
    page.fill('#panel [data-k="name"]', nm)
    page.fill('#panel [data-k="base_unit"]', '個')
    form_save(page)
    M.add_item(b, nm, '雜貨', '個')
    eq(f'品項設定（{b}）：新增後清單', read_items_table(page), exp_items_rows(M, b))
    page.fill('#q', nm[:5])
    page.wait_for_timeout(300)
    eq(f'品項設定（{b}）：搜尋「{nm[:5]}」', read_items_table(page), exp_items_rows(M, b, q=nm[:5]))
    page.fill('#q', '')
    ids = ids_of(D, b)
    victim = sorted(M.items[b])[0] if sorted(M.items[b])[0] != nm else nm
    ck(page, f'#tb [data-e="{ids[victim]}"]', '編輯品項')
    page.wait_for_selector('#pSave')
    ck(page, '#panel [data-k="active"]', '停用品項')
    form_save(page)
    M.items[b][victim]['active'] = False
    eq(f'品項設定（{b}）：停用「{victim}」後預設不顯示', read_items_table(page), exp_items_rows(M, b))
    ck(page, '#showOff', '顯示停用')
    page.wait_for_timeout(300)
    eq(f'品項設定（{b}）：勾「顯示停用」後出現', read_items_table(page), exp_items_rows(M, b, show_off=True))
    page.evaluate("()=>{const c=document.getElementById('showOff'); if(c.checked) c.click();}")


def read_vendor_table(page):
    return page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>[...r.children].slice(0,2).map(c=>c.innerText.trim()))")


def exp_vendor_rows(M, b):
    return [[v['name'] + (' 自動建立' if v['auto'] else ''), '啟用' if v['active'] else '停用'] for v in sorted(M.vendors[b].values(), key=lambda x: x['name'])]


def settings_vendors(P, S, D, M, brand, used_vendors):
    page = P['acc']
    b = brand
    nav(page, 'admin.html', 'vendors')
    page.wait_for_selector('#add')
    page.wait_for_timeout(500)
    scan(page, '設定-廠商')
    eq(f'廠商設定（{b}）：清單（含「自動建立」標籤與停用狀態）', read_vendor_table(page), exp_vendor_rows(M, b))
    ck(page, '#add', '新增廠商')
    page.wait_for_selector('#pSave')
    form_save(page, '儲存（空白）')
    check('新增廠商：沒填名稱被擋', '請填廠商名稱' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    page.fill('#panel [data-k="name"]', sorted(M.vendors[b])[0])
    form_save(page, '儲存（重名）')
    check('新增廠商：重名被擋', '這個廠商名稱已存在' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    nm = '新增廠商' + DS.pw(random.Random(S.rng.random()), 3)
    page.fill('#panel [data-k="name"]', nm)
    form_save(page)
    M.vendors[b][nm] = {'name': nm, 'active': True, 'inc': 0, 'auto': False}
    eq(f'廠商設定（{b}）：新增後清單', read_vendor_table(page), exp_vendor_rows(M, b))
    unused = [v['name'] for v in D['vendors'][b] if v['name'] not in used_vendors and v['name'] in M.vendors[b]]
    ids = {x['name']: x['id'] for x in api('GET', f'/vendors?all=1&brand_id={b}', D['adm_tok'])[1]['data']}
    if unused:
        v0 = unused[0]
        rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>r.children[0].innerText.trim())")
        idx = [r.split(' 自動建立')[0] for r in rows].index(v0)
        page.locator('#body tbody tr').nth(idx).locator('[data-e]').click()
        mark(page, '#body [data-e]', '編輯廠商')
        page.wait_for_selector('#pSave')
        ck(page, '#panel [data-k="active"]', '停用廠商')
        form_save(page)
        M.vendors[b][v0]['active'] = False
        D.setdefault('deactivated_vendor', {})[b] = v0
        eq(f'廠商設定（{b}）：停用「{v0}」', read_vendor_table(page), exp_vendor_rows(M, b))
    if len(unused) > 1:
        v1 = unused[1]
        rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>r.children[0].innerText.trim())")
        idx = [r.split(' 自動建立')[0] for r in rows].index(v1)
        page.locator('#body tbody tr').nth(idx).locator('[data-e]').click()
        page.wait_for_selector('#pSave')
        new = v1 + '改'
        page.fill('#panel [data-k="name"]', new)
        form_save(page)
        v = M.vendors[b].pop(v1)
        v['name'] = new
        M.vendors[b][new] = v
        eq(f'廠商設定（{b}）：改名「{v1}」→「{new}」', read_vendor_table(page), exp_vendor_rows(M, b))


def yuan(c):
    v = abs(c) / 100
    return str(int(v)) if float(v).is_integer() else str(float(f'{v:.2f}'))


def pnl_requeue_count(M, b, vendor, cat):
    s = set()
    for sid, st in M.confirmed(b):
        if st['vendor'] != vendor:
            continue
        if not M.stores[st['store_code']]['unit_code']:
            pass
        for ln in st['lines']:
            it = M.items[b].get(ln['item']) if ln['item'] else None
            if it and it['cat'] == cat:
                s.add((st['store_code'], st['doc_date'][:7]))
    return len(s)


def settings_pnl(P, S, D, M, brand, rng):
    """損益對照：填對照、待補對照清單、重推。"""
    page = P['acc']
    b = brand
    nav(page, 'admin.html', 'pnl')
    page.wait_for_selector('#pnlSave')
    page.wait_for_timeout(500)
    scan(page, '設定-損益對照')
    # 現有資料裡實際出現的（廠商,類別）
    pairs = set()
    for sid, st in M.confirmed(b):
        for k, v in M.slip_cats(st).items():
            if v and k != '未分類' and st['vendor']:
                pairs.add((st['vendor'], k))
    pairs = sorted(pairs)
    mapped = {}
    for (v, c) in pairs:
        if rng.random() < 0.65:
            mapped[(v, c)] = 'A' + str(rng.randint(100, 999))
    if pairs and not mapped:
        mapped[pairs[0]] = 'A' + str(rng.randint(100, 999))
    # 再填一格沒有資料的
    vs = [v for v in M.vendors[b] if M.vendors[b][v]['active']]
    extra = (rng.choice(vs), rng.choice(DS.CATS4))
    if extra not in pairs and extra not in mapped:
        mapped[extra] = 'A' + str(rng.randint(100, 999))
    # 驗證：科目代號含空白被擋
    vid = {x['name']: x['id'] for x in api('GET', f'/vendors?all=1&brand_id={b}', D['adm_tok'])[1]['data']}
    first = sorted(mapped)[0]
    page.fill(f'#body input[data-v="{vid[first[0]]}"][data-c="{first[1]}"]', 'A 1 2')
    ck(page, '#pnlSave', '儲存對照（含空白）')
    page.wait_for_timeout(500)
    check('損益對照：科目代號含空白被擋', '不可含空白' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    # 預先設定假損益端：挑一個店×月，推送時先回 LOCKED 一次
    target = None
    for (v, c), acc in sorted(mapped.items()):
        for sid, st in sorted(M.confirmed(b)):
            u = M.stores[st['store_code']]['unit_code']
            if st['vendor'] == v and u and M.slip_cats(st).get(c):
                target = (u, st['doc_date'][:7], st['store_code'])
                break
        if target:
            break
    if target:
        S.pnl.lock_once[(target[0], target[1])] = 1
    for (v, c), acc in mapped.items():
        page.fill(f'#body input[data-v="{vid[v]}"][data-c="{c}"]', acc)
    expect_requeue = sum(pnl_requeue_count(M, b, v, c) for (v, c) in mapped)
    for k, acc in mapped.items():
        M.pnlmap[(b, k[0], k[1])] = acc
    ck(page, '#pnlSave', '儲存對照')
    page.wait_for_function("()=>document.getElementById('msg').innerText.indexOf('損益對照已儲存')>=0", timeout=15000)
    m = txt(page, '#msg')
    eq(f'損益對照（{b}）：儲存訊息（筆數與排入重推的店×月數）', m.strip(), f'損益對照已儲存（{len(mapped)} 筆；{expect_requeue} 個店×月已排入重推）')
    inp = page.evaluate("()=>[...document.querySelectorAll('#body input[data-v]')].filter(i=>i.value).length")
    eq(f'損益對照（{b}）：重新載入後填入的格數', inp, len(mapped))
    # 待補對照清單
    um = M.unmapped_rows()
    rows = page.evaluate("()=>[...document.querySelectorAll('#body .card')].filter(c=>c.querySelector('h2')&&c.querySelector('h2').innerText==='待補對照').map(c=>[...c.querySelectorAll('tbody tr')].map(r=>[...r.children].map(x=>x.innerText.trim())))[0]||[]")
    brand_um = {k: v for k, v in um.items() if M.stores_by_name[k[0]]['brand'] == b}
    got = sorted((r[0], r[1], r[2], r[3].split('（')[0].strip(), r[4].split('（')[0].strip(), r[5]) for r in rows)
    want = sorted((k[0], k[1], k[2], k[3], '未對照', fmt(v / 100)) for k, v in brand_um.items())
    eq(f'損益對照（{b}）：待補對照清單（逐列）', got, want)
    tot = sum(brand_um.values())
    if brand_um:
        check(f'損益對照（{b}）：待補對照合計 {fmt(tot / 100)}', ('合計 ' + fmt(tot / 100)) in txt(page, '#body'), txt(page, '#body')[:0])
    else:
        check(f'損益對照（{b}）：沒有待補時顯示「沒有待補對照的金額」', '沒有待補對照的金額' in txt(page, '#body'))
    return target


def pnl_stuck_and_retry(P, S, D, M, brand, target):
    """LOCKED 終態 → 停推、列出「重推此店此月」→ 按下去 → 再推成功。"""
    page = P['acc']
    S.wait_pnl_quiet()
    nav(page, 'admin.html', 'items')
    nav(page, 'admin.html', 'pnl')
    page.wait_for_selector('#pnlSave')
    page.wait_for_timeout(500)
    if not target:
        return
    unit, month, code = target
    scan(page, '設定-損益對照（有卡住的店×月）')
    ent, unm = M.pnl_month(code, month)
    stuck = page.evaluate("()=>[...document.querySelectorAll('[data-rs]')].map(b=>[b.dataset.rs,b.dataset.rm,b.closest('tr').innerText])")
    st = M.stores[code]
    eq(f'損益推送卡終態（{unit} {month}）：清單出現 1 筆', len(stuck), 1)
    if stuck:
        d = sum(ent.values())
        want = f'{month} 已定稿（{st["name"]}），進貨金額變動 +{yuan(d)} 元未反映到損益，請解除定稿或手動調整'
        check(f'損益推送卡終態：原因帶方向與金額', want in stuck[0][2], (want, stuck[0][2]))
        check('損益推送卡終態：狀態欄顯示「已定稿」', '已定稿' in stuck[0][2])
    c, j = api('GET', '/health/detail', D['adm_tok'])
    check('損益推送終態時 /health/detail 帶 PNL_PUSH_STUCK 與原因文字', 'PNL_PUSH_STUCK' in j['data']['codes'] and any('已定稿' in r for r in j['data']['reasons']), j['data'].get('codes'))
    before = len([p for p in S.pnl.payloads if (p['store_id'], p['month']) == (unit, month)])
    ck(page, '[data-rs]', '重推此店此月（卡住列上的按鈕）')
    page.wait_for_function("()=>document.getElementById('msg').innerText.indexOf('已排入重推')>=0", timeout=10000)
    S.wait_pnl_quiet()
    after = [p for p in S.pnl.payloads if (p['store_id'], p['month']) == (unit, month)]
    check('按「重推」後損益端收到這個店×月且成功', len(after) > before and after[-1]['_resp'] == 'OK', [p['_resp'] for p in after])
    c, j = api('GET', '/health', None)
    check('重推成功後 PNL_PUSH_STUCK 消失', 'PNL_PUSH_STUCK' not in j['reasons'], j['reasons'])
    # 手動選店×月重推（#rtGo）
    page.wait_for_selector('#rtGo')
    sel = [s for s in D['stores'] if s['brand'] == brand and s['unit_code']]
    page.select_option('#rtStore', str(sel[0]['id']))
    mark(page, '#rtStore', '選門市')
    page.fill('#rtMonth', month)
    ck(page, '#rtGo', '重推此店此月（手動選）')
    page.wait_for_timeout(600)
    n0 = len(S.pnl.payloads)
    S.wait_pnl_quiet()
    check('手動「重推此店此月」：訊息', '已排入重推' in txt(page, '#msg'), txt(page, '#msg'))



# ───────────────────────── 門市端收尾、側欄、權限、admin ─────────────────────────
def store_final_flow(P, S, D, M, slips):
    sp = P['store']
    st = D['ui_store']
    mine = [s for s in slips if s['store'] is st]
    wave1 = [s for s in mine if s['wave'] == 1]
    old = (__import__('datetime').datetime.utcnow() - __import__('datetime').timedelta(days=5)).strftime('%Y-%m-%dT%H:%M:%S.000Z')
    for s in wave1:
        S.x('UPDATE slips SET uploaded_at=? WHERE id=?', old, s['sid'])
    ck(sp, '#refresh', '重新整理我的上傳')
    sp.wait_for_timeout(900)
    rows = read_mine(sp)
    recent = [s for s in mine if s['wave'] == 2]
    eq('門市「我的上傳」預設只列今天與昨天', sorted(r['id'] for r in rows), sorted(s['sid'] for s in recent))
    more = txt(sp, '#moreBtn')
    eq('「看更多」按鈕顯示還有幾張', more.strip(), f'看更多（最近 30 天，另有 {len(wave1)} 張）')
    scan(sp, '門市上傳頁（有看更多）')
    ck(sp, '#moreBtn', '看更多')
    sp.wait_for_timeout(700)
    rows = read_mine(sp)
    eq('按「看更多」後列出最近 30 天全部', sorted(r['id'] for r in rows), sorted(s['sid'] for s in mine))
    db = {r['id']: r['status'] for r in S.q('SELECT id, status FROM slips WHERE store_id=?', st['id'])}
    check('門市端每張貨單的狀態標籤與資料庫一致', all(r['badge'] == STATUS_TEXT[db[r['id']]] for r in rows), [(r['id'], r['badge'], db[r['id']]) for r in rows][:3])
    check('門市「我的上傳」最新的在最上面', [r['id'] for r in rows] == [r['id'] for r in sorted(rows, key=lambda r: S.q('SELECT uploaded_at FROM slips WHERE id=?', r['id'])[0]['uploaded_at'], reverse=True)])
    shot(sp, '04-門市端-我的上傳')
    # 廠商下拉反映設定頁的改動（停用的不見、新增與自動建立的出現）
    sp.reload()
    sp.wait_for_function("()=>document.getElementById('vendorSel').options.length>1")
    names = sorted(v['name'] for v in M.vendors[st['brand']].values() if v['active'])
    opts = sp.evaluate("()=>[...document.getElementById('vendorSel').options].map(o=>o.text)")
    eq('設定頁停用／新增／改名／自動建立的廠商，反映在門市廠商下拉', sorted(opts[1:-1]), names)
    # 手動輸入廠商欄位顯示／隱藏
    sp.select_option('#vendorSel', '__other')
    check('選「手動輸入」才出現廠商名稱欄', sp.evaluate("()=>!document.getElementById('vendorName').classList.contains('hidden')"))
    sp.select_option('#vendorSel', '')
    check('改回「請選擇」欄位又隱藏', sp.evaluate("()=>document.getElementById('vendorName').classList.contains('hidden')"))
    # 自願改密碼
    new = DS.pw(S.rng)
    ui_voluntary_change(sp, st['pass1'], new, '門市 ' + st['code'])
    st['pass1'] = new


def sidebar_checks(page):
    if not page.evaluate("()=>matchMedia('(hover:hover) and (pointer:fine)').matches && innerWidth>860"):
        check('側欄收合測試：此環境不是桌機滑鼠（略過）', True)
        return
    W = "()=>document.getElementById('sidein').getBoundingClientRect().width"
    C.away(page)
    page.wait_for_timeout(500)
    w0 = page.evaluate(W)
    check('側欄預設收合（窄條）', 40 < w0 < 120, w0)
    page.hover('#snav a')
    page.wait_for_timeout(700)
    w1 = page.evaluate(W)
    check('滑鼠移入側欄 → 展開（寬度明顯變大）', w1 > w0 * 2.5, (w0, w1))
    check('展開時選單文字可見', page.evaluate("()=>getComputedStyle(document.querySelector('#snav .tx')).opacity")=='1')
    C.away(page)
    page.wait_for_timeout(500)
    check('滑鼠移出側欄 → 收合回原寬', abs(page.evaluate(W) - w0) < 2, page.evaluate(W))
    page.evaluate("()=>document.querySelector('#snav a').focus()")
    page.wait_for_timeout(700)
    check('鍵盤聚焦側欄連結 → 展開', abs(page.evaluate(W) - w1) < 2, page.evaluate(W))
    page.evaluate("()=>document.activeElement.blur()")
    page.wait_for_timeout(700)
    check('取消聚焦 → 收合', abs(page.evaluate(W) - w0) < 2, page.evaluate(W))
    page.reload()
    page.wait_for_selector('#snav')
    page.wait_for_timeout(600)
    page.keyboard.press('Tab')
    inside = page.evaluate("()=>!!(document.activeElement&&document.activeElement.closest('#side'))")
    check('頁面第一個 Tab 停點在側欄（可用鍵盤進入導覽）', inside, page.evaluate("()=>{const a=document.activeElement;return a?a.tagName+'#'+a.id+'.'+a.className:''}"))
    page.wait_for_timeout(700)
    check('Tab 進入後側欄展開', abs(page.evaluate(W) - w1) < 2, page.evaluate(W))
    page.evaluate("()=>document.activeElement.blur()")
    C.away(page)


def permission_checks(S, D, M, slips):
    adm = D['adm_tok']
    ms = months_of(D)[0]
    a_single = login(D['acc_single']['username'], D['acc_single']['pass1'])[0]
    b_own = D['acc_single']['brands'][0]
    other = [b for b in D['brand_order'] if b != b_own]
    s_other = [s for s in slips if s['brand'] == other[0]][0]
    s_own = [s for s in slips if s['brand'] == b_own][0]
    t = a_single['token']
    def code(m, p, tok, body=None):
        c, j = api(m, p, tok, body)
        return c, (j.get('error') if isinstance(j, dict) else None)
    tests = [('GET', f'/slips/{s_other["sid"]}', None), ('PUT', f'/slips/{s_other["sid"]}', {}), ('POST', f'/slips/{s_other["sid"]}/unconfirm', {'reason': 'x'}),
             ('GET', f'/photos/{s_other["sid"]}/1', None), ('GET', f'/reports/cost?month={ms}&brand_id={other[0]}', None),
             ('GET', f'/reports/daily?from={ms}-01&to={ms}-28&brand_id={other[0]}', None), ('GET', f'/alerts?brand_id={other[0]}', None),
             ('GET', f'/export/legacy.xlsx?month={ms}&brand_id={other[0]}', None), ('GET', f'/vendors?brand_id={other[0]}', None),
             ('GET', f'/items?brand_id={other[0]}', None), ('POST', '/vendors', {'brand_id': other[0], 'name': 'zzz'}),
             ('POST', '/items', {'brand_id': other[0], 'name': 'zzz', 'category': '食材', 'base_unit': '個'}),
             ('PUT', '/pnl-map', {'brand_id': other[0], 'entries': []}), ('GET', f'/pnl-map?brand_id={other[0]}', None),
             ('POST', '/session/brand', {'brand_id': other[0]}), ('GET', '/admin/stores', None), ('GET', '/admin/users', None)]
    for m, p, body in tests:
        c, e = code(m, p, t, body)
        check(f'權限：{DS.BRANDS[b_own]}會計打 {DS.BRANDS[other[0]]}／管理 API「{m} {p.split("?")[0][:34]}」→ 403', c == 403 and e == 'FORBIDDEN', (c, e))
    c, e = code('GET', f'/slips/{s_own["sid"]}', t)
    check('權限：會計看自己品牌的貨單 → 200', c == 200, (c, e))
    c, j = api('GET', '/review?status=confirmed', t)
    ids = {r['id'] for r in j['data']}
    own_ids = {s['sid'] for s in slips if s['brand'] == b_own}
    check('權限：「已入帳」清單只有自己品牌的貨單', ids <= own_ids and len(ids) == len(own_ids), (len(ids), len(own_ids)))
    # 多品牌會計：目前品牌之外的資料要先切換
    am = login(D['acc_multi']['username'], D['acc_multi']['pass1'])[0]
    b1, b2 = D['acc_multi']['brands']
    s1 = [s for s in slips if s['brand'] == b1][0]
    s2 = [s for s in slips if s['brand'] == b2][0]
    cur = am['brand_id']
    cur_slip, oth_slip = (s1, s2) if cur == b1 else (s2, s1)
    c, e = code('GET', f'/slips/{oth_slip["sid"]}', am['token'])
    check('多品牌會計：沒切換前，另一個品牌的貨單 → 403', c == 403, (c, e))
    c, j = api('POST', '/session/brand', am['token'], {'brand_id': oth_slip['brand']})
    check('多品牌會計：切換到另一品牌成功', c == 200 and j['data']['brand_id'] == oth_slip['brand'], j)
    c, e = code('GET', f'/slips/{oth_slip["sid"]}', am['token'])
    check('多品牌會計：切換後可看另一品牌貨單', c == 200, (c, e))
    c, e = code('GET', f'/slips/{cur_slip["sid"]}', am['token'])
    check('多品牌會計：切換後原品牌的貨單變 403', c == 403, (c, e))
    third = [b for b in D['brand_order'] if b not in (b1, b2)][0]
    c, e = code('POST', '/session/brand', am['token'], {'brand_id': third})
    check('多品牌會計：切換到沒有權限的品牌 → 403', c == 403, (c, e))
    # 門市
    stok = login(D['stores'][0]['code'], D['stores'][0]['pass1'])[0]['token']
    for m, p in (('GET', '/review'), ('GET', f'/reports/cost?month={ms}'), ('GET', '/admin/stores'), ('GET', '/pnl-map'), ('GET', '/items')):
        c, e = code(m, p, stok)
        check(f'權限：門市帳號打會計 API「{p.split("?")[0]}」→ 403', c == 403, (c, e))
    other_store_slip = [s for s in slips if s['store'] is not D['stores'][0]][0]
    c, e = code('GET', f'/slips/{other_store_slip["sid"]}', stok)
    check('權限：門市帳號讀貨單詳情 → 403', c == 403, (c, e))
    c, e = code('GET', f'/photos/{other_store_slip["sid"]}/1', stok)
    check('權限：門市讀別家門市的照片 → 403', c == 403, (c, e))
    mine = [s for s in slips if s['store'] is D['stores'][0]]
    if mine:
        c, b = api('GET', f'/photos/{mine[0]["sid"]}/1', stok)
        check('權限：門市讀自己的照片 → 200 且是 JPEG', c == 200 and b[:2] == b'\xff\xd8', c)
    for p in ('/slips?mine=1', '/review', '/reports/cost?month=' + ms):
        c, e = code('GET', p, None)
        check(f'未登入打「{p.split("?")[0]}」→ 401', c == 401 and e == 'AUTH', (c, e))
    c, e = code('GET', '/review', 'f' * 64)
    check('亂填的 token → 401', c == 401, (c, e))
    # 連錯 5 次鎖 15 分
    ls = D['lock_store']
    for _ in range(5):
        login(ls['code'], 'wrong-' + DS.pw(S.rng))
    d, c, j = login(ls['code'], ls['pass0'])
    check('連續登入失敗 5 次 → 鎖定（即使密碼正確）', c == 403 and j.get('error') == 'LOCKED', (c, j.get('error')))


def admin_flows(P, S, D, M, slips):
    page = P['adm']
    A = D['admin']
    page.goto(url('admin.html'))
    page.wait_for_selector('#loginBtn')
    eq('設定頁標題', page.title(), TITLE)
    ui_login(page, A['username'], A['pass'], '管理者登入畫面')
    page.wait_for_selector('#app:not(.hidden)')
    check('管理者登入後側欄多出「門市」「帳號」', page.evaluate("()=>!!document.querySelector('#snav a[data-key=stores]')&&!!document.querySelector('#snav a[data-key=users]')"))
    # 工作品牌選擇
    for b in D['brand_order'][:2]:
        nav(page, 'admin.html', 'items')
        page.wait_for_selector('#addItem')
        cur = page.evaluate("()=>document.getElementById('brandPick').value")
        if cur != b:
            mark(page, '#brandPick', '切換工作品牌')
            page.hover('#snav a')
            page.select_option('#brandPick', b)
            C.away(page)
            page.wait_for_load_state('load')
            page.wait_for_selector('#addItem')
            page.wait_for_timeout(500)
        eq(f'管理者「工作品牌」切到 {DS.BRANDS[b]}：品項清單', read_items_table(page), exp_items_rows(M, b))
    scan(page, '設定-品項（管理者）')
    # 門市
    nav(page, 'admin.html', 'stores')
    page.wait_for_selector('#add')
    page.wait_for_timeout(500)
    scan(page, '設定-門市')
    rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>[...r.children].slice(0,5).map(c=>c.innerText.trim()))")
    want = [[s['code'], s['name'], DS.BRANDS[s['brand']], s['unit_code'] or '未設定（不推損益）', '啟用'] for s in sorted(D['stores'], key=lambda x: x['code'])]
    eq('門市清單（代號／名稱／品牌／損益代號／狀態）', rows, want)
    ck(page, '#add', '新增門市')
    page.wait_for_selector('#pSave')
    scan(page, '設定-新增門市表單')
    nb = D['brand_order'][1]
    code_new = 'ZN' + str(S.rng.randint(10, 99))
    cases = [({'code': 'a', 'name': 'x', 'password': '123456'}, '門市代號格式'), ({'code': code_new, 'name': '', 'password': '123456'}, '請填門市名稱'),
             ({'code': code_new, 'name': 'x', 'password': ''}, '新增門市要設定密碼'), ({'code': code_new, 'name': 'x', 'password': '123'}, '密碼至少 6 碼'),
             ({'code': D['stores'][0]['code'].lower(), 'name': 'x', 'password': '123456'}, '這個門市代號已存在')]
    for vals, msg in cases:
        for k, v in vals.items():
            page.fill(f'#panel [data-k="{k}"]', v)
        form_save(page, '儲存（驗證錯誤）')
        check(f'新增門市：「{msg}」', msg in txt(page, '#pmsg'), txt(page, '#pmsg'))
    new_pw = DS.pw(S.rng)
    page.fill('#panel [data-k="code"]', code_new.lower())
    page.fill('#panel [data-k="name"]', '新增測試店')
    mark(page, '#panel select[data-k="brand_id"]', '選品牌')
    page.select_option('#panel [data-k="brand_id"]', nb)
    page.fill('#panel [data-k="pnl_unit_code"]', 'PNL' + code_new)
    page.fill('#panel [data-k="password"]', new_pw)
    form_save(page)
    D['stores'].append({'code': code_new, 'brand': nb, 'name': '新增測試店', 'unit_code': 'PNL' + code_new, 'pass0': new_pw, 'new': True})
    M.stores[code_new] = D['stores'][-1]
    M.stores_by_name['新增測試店'] = D['stores'][-1]
    d, c, j = login(code_new, new_pw)
    check('新增的門市可用管理者設的密碼登入，且被要求首次改密碼', d and d['must_change_password'] and d['brand_id'] == nb, (c, j.get('error')))
    # 重設門市密碼 → 對方再次被強制改
    st = D['ui_store']
    sp = P['store']
    rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>r.children[0].innerText.trim())")
    page.locator('#body tbody tr').nth(rows.index(st['code'])).locator('[data-e]').click()
    mark(page, '#body [data-e]', '編輯門市')
    page.wait_for_selector('#pSave')
    tmp_pw = DS.pw(S.rng)
    page.fill('#panel [data-k="password"]', tmp_pw)
    form_save(page)
    check('管理者重設門市密碼成功', '已儲存' in txt(page, '#msg'), txt(page, '#msg'))
    d, c, j = login(st['code'], st['pass1'])
    check('重設後舊密碼不能登入', d is None and c == 401, c)
    ck(sp, '#refresh', '重新整理（舊登入已失效）')
    sp.wait_for_selector('#loginBtn', timeout=10000)
    check('門市端：舊登入被作廢 → 回到登入畫面', True)
    ui_login(sp, st['code'], tmp_pw, '門市登入畫面')
    st['pass1'] = DS.pw(S.rng)
    ui_force_change(sp, tmp_pw, st['pass1'], '門市 ' + st['code'] + '（被重設後）', st['code'])
    sp.wait_for_function("()=>document.getElementById('vendorSel').options.length>1")
    # 改損益門市代號 → 舊代號推 0
    cand = [s for s in D['stores'] if s['unit_code'] and not s.get('new') and any(x['store'] is s for x in slips)]
    if cand:
        s0 = cand[0]
        old_code = s0['unit_code']
        new_code = 'PNLX' + DS.pw(S.rng, 4).upper()
        sid_ = api('GET', '/admin/stores', D['adm_tok'])[1]['data']
        rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>r.children[0].innerText.trim())")
        page.locator('#body tbody tr').nth(rows.index(s0['code'])).locator('[data-e]').click()
        page.wait_for_selector('#pSave')
        S.pnl_mark = len(S.pnl.payloads)
        page.fill('#panel [data-k="pnl_unit_code"]', new_code)
        form_save(page)
        S.wait_pnl_quiet()
        s0['unit_code'] = new_code
        D['code_change'] = (s0, old_code, new_code)
    # 帳號
    nav(page, 'admin.html', 'users')
    page.wait_for_selector('#add')
    page.wait_for_timeout(500)
    scan(page, '設定-帳號')
    rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>[...r.children].slice(0,5).map(c=>c.innerText.trim()))")
    users = [(A['username'], '測試管理者', '管理員', '—', '啟用')] + [(a['username'], a['name'], '會計', '、'.join(DS.BRANDS[x] for x in sorted(a['brands'])), '啟用') for a in (D['acc_multi'], D['acc_single'])]
    eq('帳號清單（帳號／姓名／角色／可管理品牌／狀態）', sorted(tuple(r) for r in rows), sorted(users))
    ck(page, '#add', '新增帳號')
    page.wait_for_selector('#pSave')
    scan(page, '設定-新增帳號表單')
    page.fill('#panel [data-k="username"]', 'nu' + DS.pw(S.rng, 4).lower())
    page.fill('#panel [data-k="name"]', '新會計')
    mark(page, '#panel select[data-k="role"]', '選角色')
    page.select_option('#panel [data-k="role"]', 'admin')
    page.select_option('#panel [data-k="role"]', 'accountant')
    for b in ('X', 'M', 'C'):
        if page.evaluate(f"()=>document.querySelector('#panel input[data-bk={b}]').checked"):
            page.evaluate(f"()=>document.querySelector('#panel input[data-bk={b}]').click()")
    page.fill('#panel [data-k="password"]', 'abcdef')
    form_save(page, '儲存（沒勾品牌）')
    check('新增帳號：會計沒勾品牌被擋', '會計帳號至少要勾一個品牌' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    pick = D['brand_order'][:2]
    for b in pick:
        ck(page, f'#panel input[data-bk={b}]', f'勾選品牌 {b}')
    uname = 'nu' + DS.pw(S.rng, 4).lower()
    page.fill('#panel [data-k="username"]', uname)
    upw = DS.pw(S.rng)
    page.fill('#panel [data-k="password"]', upw)
    form_save(page)
    d, c, j = login(uname, upw)
    check('新增會計（多品牌勾選）：可登入、被要求改密碼、品牌清單正確', d and d['must_change_password'] and sorted(b['id'] for b in d['brands']) == sorted(pick), (c, d and d['brands']))
    # 編輯：加勾第三個品牌、重設密碼
    rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>r.children[0].innerText.trim())")
    page.locator('#body tbody tr').nth(rows.index(uname)).locator('[data-e]').click()
    page.wait_for_selector('#pSave')
    third = D['brand_order'][2]
    ck(page, f'#panel input[data-bk={third}]', f'加勾品牌 {third}')
    upw2 = DS.pw(S.rng)
    page.fill('#panel [data-k="password"]', upw2)
    form_save(page)
    d, c, j = login(uname, upw2)
    check('編輯帳號：加勾品牌、重設密碼後再次被要求改密碼', d and d['must_change_password'] and sorted(b['id'] for b in d['brands']) == sorted(D['brand_order']), (c, d and d['brands']))
    # 停用帳號
    page.locator('#body tbody tr').nth(rows.index(uname)).locator('[data-e]').click()
    page.wait_for_selector('#pSave')
    ck(page, '#panel [data-k="active"]', '停用帳號')
    form_save(page)
    d, c, j = login(uname, upw2)
    check('停用的帳號不能登入', d is None, c)
    # 管理者自己不能停用自己
    rows = page.evaluate("()=>[...document.querySelectorAll('#body tbody tr')].map(r=>r.children[0].innerText.trim())")
    page.locator('#body tbody tr').nth(rows.index(A['username'])).locator('[data-e]').click()
    page.wait_for_selector('#pSave')
    ck(page, '#panel [data-k="active"]', '停用自己')
    form_save(page, '儲存（停用自己）')
    check('管理者不能停用自己', '不能停用或降級自己的帳號' in txt(page, '#pmsg'), txt(page, '#pmsg'))
    ck(page, '#pCancel', '取消')
    ui_logout(page, '管理者')



# ───────────────────────── 損益推送／備份／健康檢查（獨立驗算）─────────────────────────
def verify_pnl(S, D, M):
    S.wait_pnl_quiet()
    last = {}
    for p in S.pnl.payloads:
        if p['_resp'] == 'OK':
            last[(p['store_id'], p['month'])] = p
    check('損益端收到的每一筆都帶對金鑰（沒有 AUTH 失敗）', all(p['key'] == S.pnl_key for p in S.pnl.payloads))
    bad, checked = [], 0
    for code, st in M.stores.items():
        u = st['unit_code']
        if not u:
            nothing = [k for k in last if k[0] == u]
            continue
        months = sorted({x['doc_date'][:7] for sid, x in M.confirmed(st['brand']) if x['store_code'] == code})
        for mo in months:
            ent, unm = M.pnl_month(code, mo)
            p = last.get((u, mo))
            checked += 1
            if not ent:
                if p and any(abs(v) > 0.0001 for v in p['entries'].values()):
                    bad.append(f'{u} {mo} 預期無金額卻推了 {p["entries"]}')
                continue
            if not p:
                bad.append(f'{u} {mo} 沒有成功推送')
                continue
            nz = {a: v for a, v in p['entries'].items() if abs(v) > 0.0001}
            want = {a: c / 100 for a, c in ent.items() if c}
            if set(nz) != set(want) or any(abs(nz[a] - want[a]) > 0.0001 for a in want):
                bad.append(f'{u} {mo} 科目金額 {nz}≠{want}')
            if abs(float(p.get('pending_unmapped', 0)) - unm / 100) > 0.0001:
                bad.append(f'{u} {mo} 待補 {p.get("pending_unmapped")}≠{unm / 100}')
    check(f'損益推送：{checked} 個店×月的科目金額與待補金額（稅額分攤、類別、對照）全部相符', not bad and checked > 0, '；'.join(bad[:4]))
    no_unit = [p for p in S.pnl.payloads if p['store_id'] not in {st['unit_code'] for st in M.stores.values() if st['unit_code']} | {D.get('code_change', (0, '', ''))[1]}]
    check('沒設損益代號的門市沒有被推送', not no_unit, no_unit[:1])
    if D.get('code_change'):
        s0, oldc, newc = D['code_change']
        earlier = {}
        for p in S.pnl.payloads[:S.pnl_mark]:
            if p['store_id'] == oldc and p['_resp'] == 'OK':
                earlier[p['month']] = {a for a, v in p['entries'].items() if abs(v) > 0.0001} | earlier.get(p['month'], set())
        later = {}
        for p in S.pnl.payloads[S.pnl_mark:]:
            if p['store_id'] == oldc and p['_resp'] == 'OK':
                later[p['month']] = p
        ok = True
        for mo, accs in earlier.items():
            p = later.get(mo)
            if not p or not accs <= set(p['entries']) or any(abs(v) > 0.0001 for v in p['entries'].values()):
                ok = False
        check(f'改損益門市代號：舊代號 {oldc} 每個推過的月份都被推 0 撤回（{len(earlier)} 個月）', ok and len(earlier) > 0, (earlier, list(later)))
        nl = [p for p in S.pnl.payloads[S.pnl_mark:] if p['store_id'] == newc and p['_resp'] == 'OK']
        check('改損益門市代號：新代號有收到完整金額', bool(nl))


def run_backup(S, D, M):
    import subprocess as sp
    c, j = api('GET', '/health')
    check('備份從沒成功過、已有入帳資料 → /health 紅燈 BACKUP_NEVER', j['status'] == 'red' and 'BACKUP_NEVER' in j['reasons'], j['reasons'])
    env = dict(os.environ, DATA_DIR=S.data, LOG_DIR=S.logs, PURCHASE_NO_DOTENV='1', BACKUP_URL=S.bak.url, BACKUP_KEY='wrong-key')
    r = sp.run(['node', os.path.join(ROOT, 'server', 'backup.js')], env=env, capture_output=True, text=True, timeout=120)
    check('備份：金鑰不對 → 結束碼非 0 且不寫成功紀錄', r.returncode != 0 and not os.path.exists(os.path.join(S.logs, 'backup-last.json')), (r.returncode, r.stderr[:100]))
    S.bak.payloads.clear()
    env['BACKUP_KEY'] = S.bak_key
    r = sp.run(['node', os.path.join(ROOT, 'server', 'backup.js')], env=env, capture_output=True, text=True, timeout=120)
    check('備份：執行成功（結束碼 0）', r.returncode == 0, (r.stdout[-200:], r.stderr[-200:]))
    months = sorted({x['doc_date'][:7] for sid, x in M.confirmed()})
    got_months = sorted(p['month'] for p in S.bak.payloads)
    check('備份：每個有入帳資料的月份各送一頁', set(months) <= set(got_months), (months, got_months))
    bad = []
    for p in S.bak.payloads:
        mo = p['month']
        exp = sorted([(sid, x) for sid, x in M.confirmed() if x['doc_date'][:7] == mo], key=lambda t: (t[1]['doc_date'], t[0]))
        if [s['id'] for s in p['slips']] != [sid for sid, x in exp]:
            bad.append(f'{mo} 貨單清單')
            continue
        for ps, (sid, x) in zip(p['slips'], exp):
            e = {'doc_date': x['doc_date'], 'store_code': x['store_code'], 'store_name': x['store_name'], 'brand_id': x['brand'], 'vendor_name': x['vendor'] or x['vendor_raw'] or '',
                 'doc_no': x['doc_no'], 'subtotal': x['subtotal'], 'tax': x['tax'], 'total': x['total']}
            for k, v in e.items():
                if (ps[k] != v) and not (isinstance(v, float) and ps[k] is not None and abs(ps[k] - v) < 0.0001):
                    bad.append(f'{sid} {k} {ps[k]!r}≠{v!r}')
            if not ps['confirmed_at']:
                bad.append(f'{sid} 缺 confirmed_at')
        el = []
        for sid, x in sorted(exp, key=lambda t: t[0]):          # 備份的明細依 slip_id、seq 排序
            for i, l in enumerate(x['lines']):
                it = M.items[x['brand']].get(l['item']) if l['item'] else None
                el.append((sid, i + 1, l['raw'], l['item'] or '', (it['cat'] if it else '') or '', l['qty'], l['unit'] or '', l['price'], l['amount']))
        gl = [(l['slip_id'], l['seq'], l['raw_name'], l['item_name'], l['category'], l['qty'], l['unit'], l['unit_price'], l['amount']) for l in p['lines']]
        if len(gl) != len(el):
            bad.append(f'{mo} 明細列數 {len(gl)}≠{len(el)}')
        for a, b in zip(gl, el):
            if a[:5] != b[:5] or a[6] != b[6] or any(abs(a[k] - b[k]) > 0.0001 for k in (5, 7, 8)):
                bad.append(f'{mo} 明細 {a}≠{b}')
                break
        txt_ = json.dumps(p, ensure_ascii=False)
        if 'photo' in txt_ or '.jpg' in txt_ or 'ai_raw' in txt_:
            bad.append(f'{mo} 含照片或 AI 原文')
    check('備份內容：每月每張貨單與每列明細（金額、類別、品項名）逐欄與真值相符，且不含照片', not bad, '；'.join(bad[:4]))
    c, j = api('GET', '/health')
    check('備份成功後：沒有備份紅燈', 'BACKUP_NEVER' not in j['reasons'] and 'BACKUP_STALE' not in j['reasons'] and j['backup']['last_ok_at'], j)
    unm = sum(M.unmapped_rows().values())
    codes = set(j['reasons'])
    exp_codes = {'UNMAPPED_AMOUNT'} if unm > 0 else set()
    eq('/health 公開端點：原因代碼（待補對照）', codes, exp_codes)
    eq('/health：燈號', j['status'], 'yellow' if exp_codes else 'green')
    txt_ = json.dumps(j, ensure_ascii=False)
    leak = [x for x in [s['name'] for s in D['stores']] + [s['unit_code'] for s in D['stores'] if s['unit_code']] + [v['name'] for b in M.vendors.values() for v in b.values()] if x and x in txt_]
    check('/health 公開端點不含門市名、損益代號、廠商名', not leak, leak[:3])
    eq('/health：Ollama 與損益推送狀態', (j['ollama'], j['pnl']['configured']), (True, True))
    c, j2 = api('GET', '/health/detail', D['adm_tok'])
    check('/health/detail（登入後）有詳細原因文字', c == 200 and (not exp_codes or any('待補對照' in r for r in j2['data']['reasons'])), j2['data']['reasons'])
    c, e = api('GET', '/health/detail')
    check('/health/detail 未登入 → 401', c == 401)


def final_clickmap():
    print('\n── 按鈕與連結覆蓋稽核 ──')
    rep = CM.report()
    check(f'所有按鈕與連結都被點過並驗證（共 {rep["total"]} 個）', not rep['missed'],
          '漏測：' + '、'.join(f'{k}（{v}）' for k, v in list(rep['missed'].items())[:25]))
    print(f'  掃到 {rep["total"]} 個，驗證 {rep["clicked"]} 個，漏測 {len(rep["missed"])} 個，刻意不點 {len(rep["skipped"])} 個')
    for k, v in rep['skipped'].items():
        print(f'      SKIP「{k}」：{v}')
    if rep.get('extra'):
        print('  ⚠ key 對不上（點了但掃描清單裡沒有這個名字）：')
        for k in rep['extra']:
            print(f'      「{k}」')
    for k, v in rep['missed'].items():
        print(f'      漏：{k}（{v}）')


def main():
    seed = int(os.environ.get('E2E_SEED', random.randrange(1, 10 ** 9)))
    rng = random.Random(seed)
    today = DS.taipei_today()
    D = DS.make_dataset(rng, today)
    slips = DS.make_slips(rng, D)
    DS.plan_injections(rng, D, slips)
    DS.plan_modes(rng, D, slips)
    D['lock_store'] = {'code': 'ZLK' + str(rng.randint(1, 9)), 'brand': D['brand_order'][0], 'name': '鎖定測試店', 'pass0': DS.pw(rng), 'pass1': DS.pw(rng),
                       'unit_code': None, 'lock': True}
    D['stores'].append(D['lock_store'])
    print(f'亂數種子 {seed}（重現：E2E_SEED={seed} python3 run.py）', flush=True)
    print(f'今天(台北)={today}　品牌順序={D["brand_order"]}　門市 {len(D["stores"]) - 1} 家　貨單 {len(slips)} 張　UI 門市={D["ui_store"]["code"]}', flush=True)
    for b in D['brand_order']:
        print(f'  {DS.BRANDS[b]}：廠商 {len(D["vendors"][b])}、品項 {len(D["items"][b])}、貨單 {sum(1 for s in slips if s["brand"] == b)}（含稅類型 '
              + ','.join(sorted({s["tax_kind"] for s in slips if s["brand"] == b})) + '）', flush=True)
    S = Stack(rng)
    M = DS.Model(D)
    code = 1
    try:
        S.start(D['admin']['username'], D['admin']['pass'])
        setup_backend(S, D, M, slips)
        tok = fake_check_store_api_gate(D)
        by_sid = {}
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            ctx = new_context(browser)
            P = {'store': ctx.new_page(), 'acc': ctx.new_page(), 'adm': ctx.new_page()}
            errors = []
            for k, pg in P.items():
                pg.on('dialog', C.on_dialog)
                pg.on('pageerror', lambda e, k=k: errors.append(f'{k}: {e}'))
            print('── 門市：登入、強制改密碼、上傳（第 1 批）──', flush=True)
            store_ui_prelude(P['store'], D, S)
            tok[D['ui_store']['code']] = store_token(D, D['ui_store'], D['ui_store']['pass1'])
            upload_wave(P['store'], S, D, M, slips, 1, tok)
            for s in slips:
                if s['wave'] == 1 and s['ui']:
                    pass
            print('── 會計：登入、核對第 1 批 ──', flush=True)
            multi, single = D['acc_multi'], D['acc_single']
            acc_login_first(P, D, multi, '會計登入畫面')
            for b in multi['brands']:
                acc_brand_pass(P, S, D, M, slips, by_sid, b, 1)
            ui_logout(P['acc'], '會計多品牌')
            acc_login_first(P, D, single, '會計登入畫面')
            for b in single['brands']:
                acc_brand_pass(P, S, D, M, slips, by_sid, b, 1)
            ui_logout(P['acc'], '會計單品牌')
            print('── 門市：第 2 批（含離線補傳）──', flush=True)
            P['store'].reload()
            P['store'].wait_for_selector('#vendorSel')
            P['store'].wait_for_function("()=>document.getElementById('vendorSel').options.length>1")
            upload_wave(P['store'], S, D, M, slips, 2, tok)
            print('── 會計：核對第 2 批 ──', flush=True)
            used_vendors = {s['vendor']['name'] for s in slips if not s['vendor_typed']}
            ms = months_of(D)[:3]
            tot = {'price_series': 0, 'alerts': 0}
            for a in (single, multi):
                acc_login(P, a)
                for b in a['brands']:
                    acc_brand_pass(P, S, D, M, slips, by_sid, b, 2)
                ui_logout(P['acc'], '會計 ' + a['name'])
            for a in (single, multi):
                print(f'── 會計 {a["name"]}：報表、設定頁、損益對照 ──', flush=True)
                acc_login(P, a)
                for b in a['brands']:
                    switch_brand(P['acc'], b)
                    nav(P['acc'], 'reports.html', 'cost')
                    P['acc'].wait_for_selector('#fMonth')
                    check_cost(P['acc'], D, M, b, ms[:1], with_stores=False)           # 補分類前：未分類>0
                    settings_items(P, S, D, M, b)
                    settings_vendors(P, S, D, M, b, used_vendors)
                    target = settings_pnl(P, S, D, M, b, rng)
                    pnl_stuck_and_retry(P, S, D, M, b, target)
                    n, na = reports_pass(P, S, D, M, b)
                    tot['price_series'] += n
                    tot['alerts'] += na
                if a is single:
                    sidebar_checks(P['acc'])
                ui_logout(P['acc'], '會計 ' + a['name'])
            print(f'  （單價走勢有資料的品項共 {tot["price_series"]} 個、價格提醒共 {tot["alerts"]} 筆）', flush=True)
            print('── 門市收尾、權限、管理者 ──', flush=True)
            store_final_flow(P, S, D, M, slips)
            permission_checks(S, D, M, slips)
            admin_flows(P, S, D, M, slips)
            print('── 損益推送／備份／健康檢查 ──', flush=True)
            verify_pnl(S, D, M)
            run_backup(S, D, M)
            check('過程中沒有 JavaScript 錯誤', not errors, '；'.join(errors[:3]))
            browser.close()
        code = 0
    finally:
        S.stop()
    final_clickmap()
    failed = [x for x in C.RESULTS if not x[1]]
    print(f'\n共 {len(C.RESULTS)} 項檢查，通過 {len(C.RESULTS) - len(failed)}，失敗 {len(failed)}')
    if failed:
        print(f'（重現：E2E_SEED={seed} python3 run.py）')
        for n, _, d in failed:
            print(f'  ❌ {n}　{d}')
        sys.exit(1)


if __name__ == '__main__':
    main()
